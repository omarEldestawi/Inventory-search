# -*- coding: utf-8 -*-
"""
Automated Verification Suite for Release 1.9
Egyptian Pharma Procurement Optimizer

Verifies:
1. R1: Complete removal of standalone sheet 'أصناف بخصم أعلى من المرجح'.
2. R2: Per-warehouse above-average discounts secondary table rendered on each warehouse PO sheet.
       Strict filtering for Unit Saving = Base Price * (Warehouse Discount - Reference Weighted Average Discount) >= 4.00 EGP.
       Strict exclusion of any items with saving < 4.00 EGP.
       Exact 9-column headers: [م, اسم الصنف في أرصدة المخازن, اسم الصنف في الكتالوج, سعر الجمهور للوحدة, الخصم المرجح %, خصم هذا المخزن %, أعلى خصم متاح %, المخزن صاحب أعلى خصم, فرق السعر (الوفر بالجنيه)].
3. R3: Visual price status highlighting in Column 8:
       - Soft Green fill (#E2EFDA / text #375623) when warehouse offers highest discount.
       - Soft Yellow fill (#FFF2CC / text #7F6000) when another warehouse offers higher discount with winning warehouse name and rate.
4. Clean notice message row rendered when warehouse has 0 items meeting >= 4.00 EGP.
5. Full integrity of Master Comparison, primary PO items tables, Out of Stock, Recommendations, and accounting formulas.
6. R4: Version bump across app_gui.py, installer.iss, and metadata to 1.9.
"""
import os
import re
import tempfile
import pytest
import openpyxl
import pandas as pd

from src.report_generator import ArabicExcelReportGenerator
from src.allocator import MedicineAllocator, evaluate_above_average_discounts
from app_gui import PharmaProcurementApp


class TestRelease19Verification:
    @pytest.fixture
    def v19_report_payload(self):
        """Constructs multi-warehouse test payload covering both winning and competing discount tiers."""
        master_df = pd.DataFrame([
            {
                'م': 1,
                'اسم الصنف المطلوب (في طلب الشراء)': 'بانادول اكسترا 24 قرص',
                'اسم الصنف في كشف المخزن الفائز': 'Panadol Extra 24 Tab',
                'الكمية': 20,
                'المخزن الفائز': 'مخزن السلام فارم',
                'أعلى خصم %': 0.30,
                'سعر الجمهور': 50.0,
                'سعر الشراء للوحدة': 35.0,
                'إجمالي سعر الجمهور': 1000.0,
                'إجمالي تكلفة الشراء': 700.0,
                'قيمة الوفر': 300.0,
                'الخصم المرجح %': 0.20,
                'فرق الخصم %': 0.10,
                'خصم مخزن السلام فارم': 0.30,
                'خصم مخزن الفرسان بنها': 0.25,
                'خصم مخزن الأمل': 0.18,
            },
            {
                'م': 2,
                'اسم الصنف المطلوب (في طلب الشراء)': 'كونكور 5 مجم 30 قرص',
                'اسم الصنف في كشف المخزن الفائز': 'Concor 5mg 30 Tab',
                'الكمية': 15,
                'المخزن الفائز': 'مخزن الفرسان بنها',
                'أعلى خصم %': 0.28,
                'سعر الجمهور': 60.0,
                'سعر الشراء للوحدة': 43.2,
                'إجمالي سعر الجمهور': 900.0,
                'إجمالي تكلفة الشراء': 648.0,
                'قيمة الوفر': 252.0,
                'الخصم المرجح %': 0.15,
                'فرق الخصم %': 0.13,
                'خصم مخزن السلام فارم': 0.22,
                'خصم مخزن الفرسان بنها': 0.28,
                'خصم مخزن الأمل': 0.16,
            }
        ])

        # Above average reference items:
        # Item 1: High saving in Salam (50 * (0.30 - 0.20) = 5.00 EGP >= 4.00) -> Salam is winner (30%)
        #         Medium saving in Forsan (50 * (0.28 - 0.20) = 4.00 EGP >= 4.00) -> Forsan is NOT winner (Salam 30.0%)
        # Item 2: High saving in Forsan (100 * (0.25 - 0.15) = 10.00 EGP >= 4.00) -> Forsan is winner (25%)
        #         Low saving in Salam (100 * (0.18 - 0.15) = 3.00 EGP < 4.00) -> Salam MUST strictly EXCLUDE!
        # Item 3: Low saving everywhere (40 * (0.25 - 0.20) = 2.00 EGP < 4.00) -> Strictly excluded everywhere!
        above_df = pd.DataFrame([
            {
                'م': 1,
                'اسم الصنف في أرصده المخازن': 'بانادول اكسترا 24 قرص',
                'اسم الصنف في كشف المخزن الفائز': 'Panadol Extra 24 Tab',
                'المخزن الفائز': 'مخزن السلام فارم',
                'الخصم المرجح %': 0.20,
                'أعلى خصم متاح %': 0.30,
                'فرق الخصم %': 0.10,
                'سعر الجمهور للوحدة': 50.0,
                'سعر التوريد الصافي': 35.0,
                'خصم مخزن السلام فارم': 0.30,
                'خصم مخزن الفرسان بنها': 0.28,
                'خصم مخزن الأمل': 0.18,
            },
            {
                'م': 2,
                'اسم الصنف في أرصده المخازن': 'كونكور 5 مجم 30 قرص',
                'اسم الصنف في كشف المخزن الفائز': 'Concor 5mg 30 Tab',
                'المخزن الفائز': 'مخزن الفرسان بنها',
                'الخصم المرجح %': 0.15,
                'أعلى خصم متاح %': 0.25,
                'فرق الخصم %': 0.10,
                'سعر الجمهور للوحدة': 100.0,
                'سعر التوريد الصافي': 75.0,
                'خصم مخزن السلام فارم': 0.18,
                'خصم مخزن الفرسان بنها': 0.25,
                'خصم مخزن الأمل': 0.16,
            },
            {
                'م': 3,
                'اسم الصنف في أرصده المخازن': 'صنف وفر ضئيل',
                'اسم الصنف في كشف المخزن الفائز': 'Low Saving Drug',
                'المخزن الفائز': 'مخزن السلام فارم',
                'الخصم المرجح %': 0.20,
                'أعلى خصم متاح %': 0.25,
                'فرق الخصم %': 0.05,
                'سعر الجمهور للوحدة': 40.0,
                'سعر التوريد الصافي': 30.0,
                'خصم مخزن السلام فارم': 0.25,
                'خصم مخزن الفرسان بنها': 0.22,
                'خصم مخزن الأمل': 0.19,
            }
        ])

        po_salam = pd.DataFrame([
            {
                'م': 1, 'اسم الصنف في أمر الشراء': 'بانادول اكسترا 24 قرص',
                'اسم الصنف في كتالوج المخزن': 'Panadol Extra 24 Tab', 'الكمية المطلوبة': 20,
                'نسبة الخصم %': 0.30, 'سعر الجمهور للوحدة': 50.0, 'سعر التوريد الصافي': 35.0,
                'إجمالي القيمة بسعر الجمهور': 1000.0, 'إجمالي القيمة الصافية': 700.0,
                'ملاحظات التوريد': 'تأكيد التوريد الفوري'
            }
        ])

        po_forsan = pd.DataFrame([
            {
                'م': 1, 'اسم الصنف في أمر الشراء': 'كونكور 5 مجم 30 قرص',
                'اسم الصنف في كتالوج المخزن': 'Concor 5mg 30 Tab', 'الكمية المطلوبة': 15,
                'نسبة الخصم %': 0.28, 'سعر الجمهور للوحدة': 60.0, 'سعر التوريد الصافي': 43.2,
                'إجمالي القيمة بسعر الجمهور': 900.0, 'إجمالي القيمة الصافية': 648.0,
                'ملاحظات التوريد': 'تأكيد التوريد الفوري'
            }
        ])

        po_amal = pd.DataFrame([
            {
                'م': 1, 'اسم الصنف في أمر الشراء': 'أوجمنتين 1 جم',
                'اسم الصنف في كتالوج المخزن': 'Augmentin 1gm Tab', 'الكمية المطلوبة': 5,
                'نسبة الخصم %': 0.15, 'سعر الجمهور للوحدة': 120.0, 'سعر التوريد الصافي': 102.0,
                'إجمالي القيمة بسعر الجمهور': 600.0, 'إجمالي القيمة الصافية': 510.0,
                'ملاحظات التوريد': 'تأكيد التوريد الفوري'
            }
        ])

        return {
            'kpis': {'total_items': 3, 'matched_items': 3, 'total_savings': 552.0},
            'master_comparison': master_df,
            'above_average_discounts': above_df,
            'allocated_orders': {
                'مخزن السلام فارم': po_salam,
                'مخزن الفرسان بنها': po_forsan,
                'مخزن الأمل': po_amal
            },
            'out_of_stock': pd.DataFrame(),
            'recommendations': pd.DataFrame()
        }

    def test_r1_standalone_sheet_strictly_removed(self, v19_report_payload):
        """R1: Standalone sheet 'أصناف بخصم أعلى من المرجح' must NOT exist in workbooks."""
        with tempfile.TemporaryDirectory() as tmp_dir:
            generator = ArabicExcelReportGenerator(v19_report_payload, output_dir=tmp_dir)
            primary_path = generator.generate(custom_date_str="2026-10-19")
            arabic_path = os.path.join(tmp_dir, "أفضل_اسعار_2026-10-19.xlsx")

            for p in [primary_path, arabic_path]:
                wb = openpyxl.load_workbook(p)
                assert "أصناف بخصم أعلى من المرجح" not in wb.sheetnames
                assert "المقارنة الشاملة وتوزيع الأصناف" in wb.sheetnames

    def test_r2_secondary_table_columns_and_strict_4egp_filtering(self, v19_report_payload):
        """
        R2: Secondary table on warehouse PO sheets:
        - Exact 9 columns.
        - Strict filtering: Unit Saving >= 4.00 EGP.
        - Items saving < 4.00 EGP are strictly excluded.
        """
        with tempfile.TemporaryDirectory() as tmp_dir:
            generator = ArabicExcelReportGenerator(v19_report_payload, output_dir=tmp_dir)
            primary_path = generator.generate(custom_date_str="2026-10-19")

            wb = openpyxl.load_workbook(primary_path)

            expected_cols = [
                'م', 'اسم الصنف في أرصدة المخازن', 'اسم الصنف في الكتالوج',
                'سعر الجمهور للوحدة', 'الخصم المرجح %', 'خصم هذا المخزن %',
                'أعلى خصم متاح %', 'المخزن صاحب أعلى خصم', 'فرق السعر (الوفر بالجنيه)'
            ]

            # 1. Inspect Salam sheet
            ws_salam = [wb[s] for s in wb.sheetnames if "السلام" in s][0]
            sec_header_row = None
            for r in range(1, ws_salam.max_row + 1):
                if ws_salam.cell(row=r, column=1).value == 'م' and ws_salam.cell(row=r, column=2).value == 'اسم الصنف في أرصدة المخازن':
                    sec_header_row = r
                    break
            assert sec_header_row is not None
            for idx, col_name in enumerate(expected_cols, 1):
                assert ws_salam.cell(row=sec_header_row, column=idx).value == col_name

            # In Salam: Item 1 has saving = 50 * (0.30 - 0.20) = 5.00 EGP >= 4.00 -> Included
            # Item 2 has saving = 100 * (0.18 - 0.15) = 3.00 EGP < 4.00 -> Strictly excluded
            # Item 3 has saving = 40 * (0.25 - 0.20) = 2.00 EGP < 4.00 -> Strictly excluded
            salam_items = []
            for r in range(sec_header_row + 1, ws_salam.max_row + 1):
                c2 = ws_salam.cell(row=r, column=2).value
                if c2 and "إجمالي" not in str(c2) and "لا توجد" not in str(c2):
                    salam_items.append((c2, ws_salam.cell(row=r, column=9).value))

            assert len(salam_items) == 1
            assert salam_items[0][0] == 'بانادول اكسترا 24 قرص'
            assert salam_items[0][1] == 5.00

            # 2. Inspect Forsan sheet
            ws_forsan = [wb[s] for s in wb.sheetnames if "الفرسان" in s][0]
            sec_hdr_forsan = None
            for r in range(1, ws_forsan.max_row + 1):
                if ws_forsan.cell(row=r, column=1).value == 'م' and ws_forsan.cell(row=r, column=2).value == 'اسم الصنف في أرصدة المخازن':
                    sec_hdr_forsan = r
                    break
            assert sec_hdr_forsan is not None

            # In Forsan:
            # Item 1 has saving = 50 * (0.28 - 0.20) = 4.00 EGP >= 4.00 -> Included (not winner)
            # Item 2 has saving = 100 * (0.25 - 0.15) = 10.00 EGP >= 4.00 -> Included (winner)
            # Item 3 has saving = 40 * (0.22 - 0.20) = 0.80 EGP < 4.00 -> Excluded
            forsan_items = []
            for r in range(sec_hdr_forsan + 1, ws_forsan.max_row + 1):
                c2 = ws_forsan.cell(row=r, column=2).value
                if c2 and "إجمالي" not in str(c2) and "لا توجد" not in str(c2):
                    forsan_items.append((c2, ws_forsan.cell(row=r, column=9).value))

            assert len(forsan_items) == 2
            for item_name, unit_saving in forsan_items:
                assert unit_saving >= 4.00

    def test_r3_visual_price_status_highlighting(self, v19_report_payload):
        """
        R3: Column 8 visual status styling:
        - Soft Green (#E2EFDA / text #375623) when this warehouse offers highest discount.
        - Soft Yellow (#FFF2CC / text #7F6000) when another warehouse offers higher discount.
        """
        with tempfile.TemporaryDirectory() as tmp_dir:
            generator = ArabicExcelReportGenerator(v19_report_payload, output_dir=tmp_dir)
            primary_path = generator.generate(custom_date_str="2026-10-19")
            wb = openpyxl.load_workbook(primary_path)

            # 1. Salam sheet: Panadol Extra is winner in Salam
            ws_salam = [wb[s] for s in wb.sheetnames if "السلام" in s][0]
            sec_header_row = [r for r in range(1, ws_salam.max_row + 1) if ws_salam.cell(row=r, column=1).value == 'م' and ws_salam.cell(row=r, column=2).value == 'اسم الصنف في أرصدة المخازن'][0]

            cell_col8_salam = ws_salam.cell(row=sec_header_row + 1, column=8)
            # Winning warehouse styling: Soft Green
            assert cell_col8_salam.fill.start_color.rgb in ['E2EFDA', '00E2EFDA', 'FFE2EFDA']
            assert cell_col8_salam.font.color.rgb in ['375623', '00375623', 'FF375623']
            assert "السلام" in str(cell_col8_salam.value)

            # 2. Forsan sheet: Contains two items
            ws_forsan = [wb[s] for s in wb.sheetnames if "الفرسان" in s][0]
            sec_hdr_forsan = [r for r in range(1, ws_forsan.max_row + 1) if ws_forsan.cell(row=r, column=1).value == 'م' and ws_forsan.cell(row=r, column=2).value == 'اسم الصنف في أرصدة المخازن'][0]

            # Row 1 in Forsan is Concor 5mg (saving 10.00, winner in Forsan)
            row_concor = sec_hdr_forsan + 1
            cell_concor_status = ws_forsan.cell(row=row_concor, column=8)
            assert cell_concor_status.fill.start_color.rgb in ['E2EFDA', '00E2EFDA', 'FFE2EFDA']
            assert cell_concor_status.font.color.rgb in ['375623', '00375623', 'FF375623']

            # Row 2 in Forsan is Panadol Extra (saving 4.00, non-winner: Salam 30.0%)
            row_panadol = sec_hdr_forsan + 2
            cell_panadol_status = ws_forsan.cell(row=row_panadol, column=8)
            # Non-winning styling: Soft Yellow
            assert cell_panadol_status.fill.start_color.rgb in ['FFF2CC', '00FFF2CC', 'FFFFF2CC']
            assert cell_panadol_status.font.color.rgb in ['7F6000', '007F6000', 'FF7F6000']
            # Displays competing warehouse name and rate
            assert "السلام" in str(cell_panadol_status.value)
            assert "30" in str(cell_panadol_status.value)

    def test_zero_items_clean_notice_row_resilience(self, v19_report_payload):
        """Clean notice message row is displayed without error when a warehouse has 0 items meeting >= 4.00 EGP."""
        with tempfile.TemporaryDirectory() as tmp_dir:
            generator = ArabicExcelReportGenerator(v19_report_payload, output_dir=tmp_dir)
            primary_path = generator.generate(custom_date_str="2026-10-19")
            wb = openpyxl.load_workbook(primary_path)

            # مخزن الأمل has 0 items saving >= 4.00 EGP
            ws_amal = [wb[s] for s in wb.sheetnames if "الأمل" in s][0]
            sec_header_row = [r for r in range(1, ws_amal.max_row + 1) if ws_amal.cell(row=r, column=1).value == 'م' and ws_amal.cell(row=r, column=2).value == 'اسم الصنف في أرصدة المخازن'][0]

            notice_cell = ws_amal.cell(row=sec_header_row + 1, column=1)
            assert "لا توجد أصناف إضافية" in str(notice_cell.value)
            sum_cell = ws_amal.cell(row=sec_header_row + 2, column=2)
            assert "إجمالي الأصناف: 0 صنف" in str(sum_cell.value)

    def test_integrity_and_accounting_formulas_preservation(self, v19_report_payload):
        """Primary PO tables and Master Comparison strictly preserve =SUM(...) accounting formulas."""
        with tempfile.TemporaryDirectory() as tmp_dir:
            generator = ArabicExcelReportGenerator(v19_report_payload, output_dir=tmp_dir)
            primary_path = generator.generate(custom_date_str="2026-10-19")
            wb = openpyxl.load_workbook(primary_path)

            # Master Comparison formula check
            ws_master = wb["المقارنة الشاملة وتوزيع الأصناف"]
            sum_row_master = 9 + len(v19_report_payload['master_comparison'])  # headers at row 8, data 9-10, sum at 11
            cell_formula_d = str(ws_master.cell(row=sum_row_master, column=4).value)
            assert "=SUM(" in cell_formula_d

            # Primary PO sheets formula check
            for sname in wb.sheetnames:
                if "طلب شراء" in sname:
                    ws = wb[sname]
                    # Row 6 is primary PO summary row (1 item in PO at row 5)
                    sum_formula_d = str(ws.cell(row=6, column=4).value)
                    sum_formula_h = str(ws.cell(row=6, column=8).value)
                    sum_formula_i = str(ws.cell(row=6, column=9).value)
                    assert "=SUM(D5:D5)" in sum_formula_d
                    assert "=SUM(H5:H5)" in sum_formula_h
                    assert "=SUM(I5:I5)" in sum_formula_i

    def test_r4_version_bump_to_19(self):
        """R4: All version strings and identifiers bump from 1.8 to 1.9."""
        # 1. Test GUI app_version and title
        app = PharmaProcurementApp()
        app.withdraw()
        try:
            assert app.app_version == "1.9"
            assert "v1.9" in app.title()
        finally:
            app.destroy()

        # 2. Test installer.iss defines MyAppVersion "1.9" and Setup_v1.9
        iss_path = os.path.join(os.path.dirname(os.path.dirname(__file__)), "installer.iss")
        with open(iss_path, "r", encoding="utf-8") as f:
            iss_content = f.read()
        assert '#define MyAppVersion "1.9"' in iss_content
        assert 'OutputBaseFilename=EgyptianPharmaOptimizer_Setup_v1.9' in iss_content

    def test_variant_catalog_entries_selects_optimal_saving(self):
        """
        Edge Case: Supplier catalog contains multiple variant entries of the same drug.
        Variant 1 has higher discount (26%) on lower base price (50.0 EGP) -> Saving: 3.00 EGP (< 4.00 EGP).
        Variant 2 has lower discount (24%) on higher base price (120.0 EGP) -> Saving: 4.80 EGP (>= 4.00 EGP).
        Ensures the qualifying variant is chosen rather than prematurely discarded.
        """
        inv = pd.DataFrame([
            {
                'warehouse_name': 'مخزن السلام فارم',
                'item_name': 'بانادول اكسترا 24 قرص',
                'normalized_name': 'بانادول اكسترا 24 قرص',
                'dosage': '', 'unit_size': '',
                'base_price': 50.0, 'discount_rate': 0.26, 'stock_status': True
            },
            {
                'warehouse_name': 'مخزن السلام فارم',
                'item_name': 'بانادول اكسترا 24 قرص',
                'normalized_name': 'بانادول اكسترا 24 قرص',
                'dosage': '', 'unit_size': '',
                'base_price': 120.0, 'discount_rate': 0.24, 'stock_status': True
            },
        ])
        ref_df = pd.DataFrame([{'clean_name': 'بانادول اكسترا 24 قرص', 'avg_discount_rate': 0.20}])
        alloc = MedicineAllocator(inv, enable_ai=False)
        res = evaluate_above_average_discounts(ref_df, alloc)
        per_wh = res.attrs.get('per_warehouse', {})
        salam_df = per_wh.get('مخزن السلام فارم', pd.DataFrame())

        assert len(salam_df) == 1, "Expected qualifying variant to be included in secondary table"
        assert salam_df.iloc[0]['سعر الجمهور للوحدة'] == 120.0
        assert salam_df.iloc[0]['فرق السعر (الوفر بالجنيه)'] == 4.80
        assert salam_df.iloc[0]['خصم هذا المخزن %'] == 0.24

    def test_empty_secondary_table_merged_notice_banner(self, v19_report_payload):
        """
        Visual layout: When a warehouse has 0 items meeting >= 4.00 EGP,
        the notice message banner is cleanly merged across columns A to I.
        """
        with tempfile.TemporaryDirectory() as tmp_dir:
            generator = ArabicExcelReportGenerator(v19_report_payload, output_dir=tmp_dir)
            primary_path = generator.generate(custom_date_str="2026-10-19")
            wb = openpyxl.load_workbook(primary_path)

            ws_amal = [wb[s] for s in wb.sheetnames if "الأمل" in s][0]
            sec_header_row = [
                r for r in range(1, ws_amal.max_row + 1)
                if ws_amal.cell(row=r, column=1).value == 'م' and ws_amal.cell(row=r, column=2).value == 'اسم الصنف في أرصدة المخازن'
            ][0]
            empty_row = sec_header_row + 1

            # Assert notice row is in merged_cells
            expected_range = f"A{empty_row}:I{empty_row}"
            merged_str_ranges = [str(rng) for rng in ws_amal.merged_cells.ranges]
            assert expected_range in merged_str_ranges, f"Expected {expected_range} to be merged, got {merged_str_ranges}"
            assert "لا توجد أصناف إضافية" in str(ws_amal.cell(row=empty_row, column=1).value)

    def test_multi_warehouse_variant_conflict_prevent_false_winner(self):
        """
        Adversarial Edge Case:
        Salam has Variant 1 (26%, 50 EGP -> saving 3.00 EGP < 4.00 EGP)
        and Variant 2 (24%, 120 EGP -> saving 4.80 EGP >= 4.00 EGP).
        Forsan has Variant 3 (25%, 100 EGP -> saving 5.00 EGP >= 4.00 EGP).

        Verifies:
        - Forsan (25%) is styled Soft Green (#E2EFDA) as winner with max discount 25.0%.
        - Salam (24%) is styled Soft Yellow (#FFF2CC) displaying Forsan (25.0%).
        - Salam (with lower discount) is NEVER falsely marked Soft Green.
        """
        inv = pd.DataFrame([
            {
                'warehouse_name': 'مخزن السلام فارم',
                'item_name': 'بانادول اكسترا 24 قرص',
                'normalized_name': 'بانادول اكسترا 24 قرص',
                'dosage': '', 'unit_size': '',
                'base_price': 50.0, 'discount_rate': 0.26, 'stock_status': True
            },
            {
                'warehouse_name': 'مخزن السلام فارم',
                'item_name': 'بانادول اكسترا 24 قرص',
                'normalized_name': 'بانادول اكسترا 24 قرص',
                'dosage': '', 'unit_size': '',
                'base_price': 120.0, 'discount_rate': 0.24, 'stock_status': True
            },
            {
                'warehouse_name': 'مخزن الفرسان بنها',
                'item_name': 'بانادول اكسترا 24 قرص',
                'normalized_name': 'بانادول اكسترا 24 قرص',
                'dosage': '', 'unit_size': '',
                'base_price': 100.0, 'discount_rate': 0.25, 'stock_status': True
            }
        ])
        ref_df = pd.DataFrame([{'clean_name': 'بانادول اكسترا 24 قرص', 'avg_discount_rate': 0.20}])
        alloc = MedicineAllocator(inv, enable_ai=False)
        res = evaluate_above_average_discounts(ref_df, alloc)

        payload = {
            'master_comparison': pd.DataFrame(),
            'above_average_discounts': res,
            'per_warehouse_above_average_discounts': res.attrs['per_warehouse'],
            'allocated_orders': {
                'مخزن السلام فارم': pd.DataFrame([{'م': 1, 'اسم الصنف في أمر الشراء': 'x', 'اسم الصنف في كتالوج المخزن': 'x', 'الكمية المطلوبة': 1, 'نسبة الخصم %': 0.2, 'سعر الجمهور للوحدة': 10, 'سعر التوريد الصافي': 8, 'إجمالي القيمة بسعر الجمهور': 10, 'إجمالي القيمة الصافية': 8, 'ملاحظات التوريد': ''}]),
                'مخزن الفرسان بنها': pd.DataFrame([{'م': 1, 'اسم الصنف في أمر الشراء': 'x', 'اسم الصنف في كتالوج المخزن': 'x', 'الكمية المطلوبة': 1, 'نسبة الخصم %': 0.2, 'سعر الجمهور للوحدة': 10, 'سعر التوريد الصافي': 8, 'إجمالي القيمة بسعر الجمهور': 10, 'إجمالي القيمة الصافية': 8, 'ملاحظات التوريد': ''}])
            }
        }

        with tempfile.TemporaryDirectory() as tmp_dir:
            gen = ArabicExcelReportGenerator(payload, output_dir=tmp_dir)
            p = gen.generate(custom_date_str='2026-10-19')
            wb = openpyxl.load_workbook(p)

            # Check Salam sheet
            ws_salam = [wb[s] for s in wb.sheetnames if "السلام" in s][0]
            sec_hdr_salam = [
                r for r in range(1, ws_salam.max_row + 1)
                if ws_salam.cell(row=r, column=1).value == 'م' and ws_salam.cell(row=r, column=2).value == 'اسم الصنف في أرصدة المخازن'
            ][0]
            data_row_salam = sec_hdr_salam + 1

            assert ws_salam.cell(row=data_row_salam, column=6).value == 0.24
            assert ws_salam.cell(row=data_row_salam, column=7).value == 0.25
            cell_status_salam = ws_salam.cell(row=data_row_salam, column=8)
            assert "الفرسان" in str(cell_status_salam.value)
            assert "25" in str(cell_status_salam.value)
            assert cell_status_salam.fill.start_color.rgb in ['FFF2CC', '00FFF2CC', 'FFFFF2CC']

            # Check Forsan sheet
            ws_forsan = [wb[s] for s in wb.sheetnames if "الفرسان" in s][0]
            sec_hdr_forsan = [
                r for r in range(1, ws_forsan.max_row + 1)
                if ws_forsan.cell(row=r, column=1).value == 'م' and ws_forsan.cell(row=r, column=2).value == 'اسم الصنف في أرصدة المخازن'
            ][0]
            data_row_forsan = sec_hdr_forsan + 1

            assert ws_forsan.cell(row=data_row_forsan, column=6).value == 0.25
            assert ws_forsan.cell(row=data_row_forsan, column=7).value == 0.25
            cell_status_forsan = ws_forsan.cell(row=data_row_forsan, column=8)
            assert "الفرسان" in str(cell_status_forsan.value)
            assert cell_status_forsan.fill.start_color.rgb in ['E2EFDA', '00E2EFDA', 'FFE2EFDA']

    def test_secondary_table_string_discount_and_currency_resilience(self):
        """
        Resilience: When discounts or prices are input as formatted strings
        (e.g. '25%', '1,250.00 ج.م', '62.50 EGP'), they are safely parsed
        and rendered as numbers, not converted to '-'.
        """
        sec_df = pd.DataFrame([{
            'م': 1,
            'اسم الصنف في أرصدة المخازن': 'دواء اختباري',
            'اسم الصنف في الكتالوج': 'Test Drug',
            'سعر الجمهور للوحدة': '1,250.00 ج.م',
            'الخصم المرجح %': '20.0%',
            'خصم هذا المخزن %': '25.0%',
            'أعلى خصم متاح %': '25.0%',
            'المخزن صاحب أعلى خصم': 'مخزن السلام فارم',
            'فرق السعر (الوفر بالجنيه)': '62.50 ج.م'
        }])
        payload = {
            'master_comparison': pd.DataFrame(),
            'per_warehouse_above_average_discounts': {'مخزن السلام فارم': sec_df},
            'allocated_orders': {
                'مخزن السلام فارم': pd.DataFrame([{'م': 1, 'اسم الصنف في أمر الشراء': 'x', 'اسم الصنف في كتالوج المخزن': 'x', 'الكمية المطلوبة': 1, 'نسبة الخصم %': 0.2, 'سعر الجمهور للوحدة': 10, 'سعر التوريد الصافي': 8, 'إجمالي القيمة بسعر الجمهور': 10, 'إجمالي القيمة الصافية': 8, 'ملاحظات التوريد': ''}])
            }
        }
        with tempfile.TemporaryDirectory() as tmp_dir:
            gen = ArabicExcelReportGenerator(payload, output_dir=tmp_dir)
            p = gen.generate(custom_date_str='2026-10-19')
            wb = openpyxl.load_workbook(p)
            ws = [wb[s] for s in wb.sheetnames if "السلام" in s][0]
            sec_hdr = [
                r for r in range(1, ws.max_row + 1)
                if ws.cell(row=r, column=1).value == 'م' and ws.cell(row=r, column=2).value == 'اسم الصنف في أرصدة المخازن'
            ][0]
            r1 = sec_hdr + 1

            assert ws.cell(row=r1, column=4).value == 1250.0
            assert ws.cell(row=r1, column=5).value == 0.2
            assert ws.cell(row=r1, column=6).value == 0.25
            assert ws.cell(row=r1, column=7).value == 0.25
            assert ws.cell(row=r1, column=9).value == 62.5
            # Soft Green styling
            assert ws.cell(row=r1, column=8).fill.start_color.rgb in ['E2EFDA', '00E2EFDA', 'FFE2EFDA']

    def test_secondary_table_accounting_sum_formula(self, v19_report_payload):
        """
        Formula Integrity: The secondary table summary row preserves
        the dynamic =SUM(I...:I...) formula for Column 9.
        """
        with tempfile.TemporaryDirectory() as tmp_dir:
            generator = ArabicExcelReportGenerator(v19_report_payload, output_dir=tmp_dir)
            primary_path = generator.generate(custom_date_str="2026-10-19")
            wb = openpyxl.load_workbook(primary_path)

            ws_salam = [wb[s] for s in wb.sheetnames if "السلام" in s][0]
            sec_hdr = [
                r for r in range(1, ws_salam.max_row + 1)
                if ws_salam.cell(row=r, column=1).value == 'م' and ws_salam.cell(row=r, column=2).value == 'اسم الصنف في أرصدة المخازن'
            ][0]
            data_row = sec_hdr + 1
            sum_row = data_row + 1

            # Assert formula is present in Column 9
            sum_cell = ws_salam.cell(row=sum_row, column=9)
            assert str(sum_cell.value).startswith("=SUM(I")
            assert f"I{data_row}" in str(sum_cell.value)

    def test_secondary_table_missing_saving_column_calculated_and_sorted(self):
        """
        Robustness: When per_warehouse_above_average_discounts does not include
        'فرق السعر (الوفر بالجنيه)', it is dynamically calculated without KeyError
        and properly filtered for >= 4.00 EGP.
        """
        sec_df = pd.DataFrame([{
            'اسم الصنف في أرصدة المخازن': 'بانادول',
            'اسم الصنف في الكتالوج': 'Panadol',
            'سعر الجمهور للوحدة': 100.0,
            'الخصم المرجح %': 0.10,
            'خصم هذا المخزن %': 0.20,
            'أعلى خصم متاح %': 0.20,
            'المخزن صاحب أعلى خصم': 'مخزن السلام فارم',
        }])
        payload = {
            'master_comparison': pd.DataFrame(),
            'per_warehouse_above_average_discounts': {'مخزن السلام فارم': sec_df},
            'allocated_orders': {
                'مخزن السلام فارم': pd.DataFrame([{'م': 1, 'اسم الصنف في أمر الشراء': 'x', 'اسم الصنف في كتالوج المخزن': 'x', 'الكمية المطلوبة': 1, 'نسبة الخصم %': 0.2, 'سعر الجمهور للوحدة': 10, 'سعر التوريد الصافي': 8, 'إجمالي القيمة بسعر الجمهور': 10, 'إجمالي القيمة الصافية': 8, 'ملاحظات التوريد': ''}])
            }
        }
        with tempfile.TemporaryDirectory() as tmp_dir:
            gen = ArabicExcelReportGenerator(payload, output_dir=tmp_dir)
            p = gen.generate(custom_date_str='2026-10-19')
            wb = openpyxl.load_workbook(p)
            ws = [wb[s] for s in wb.sheetnames if 'السلام' in s][0]
            sec_hdr = [
                r for r in range(1, ws.max_row + 1)
                if ws.cell(row=r, column=1).value == 'م' and ws.cell(row=r, column=2).value == 'اسم الصنف في أرصدة المخازن'
            ][0]
            r1 = sec_hdr + 1
            assert ws.cell(row=r1, column=2).value == 'بانادول'
            assert ws.cell(row=r1, column=9).value == 10.0
            assert ws.cell(row=r1, column=8).fill.start_color.rgb in ['E2EFDA', '00E2EFDA', 'FFE2EFDA']

    def test_secondary_table_string_formatted_saving_sorted_numerically(self):
        """
        Numeric Ordering Gate: Ensures that string savings like '5.00 ج.م' vs '10.00 ج.م'
        are parsed into numeric floats and sorted strictly descending (10.00 before 5.00),
        preventing incorrect alphabetical sorting.
        """
        cand_df = pd.DataFrame([
            {
                'اسم الصنف في أرصدة المخازن': 'صنف وفر 5',
                'اسم الصنف في الكتالوج': 'Item 5',
                'سعر الجمهور للوحدة': '100.00 ج.م',
                'الخصم المرجح %': '10.0%',
                'خصم هذا المخزن %': '15.0%',
                'أعلى خصم متاح %': '15.0%',
                'المخزن صاحب أعلى خصم': 'مخزن السلام فارم',
                'فرق السعر (الوفر بالجنيه)': '5.00 ج.م'
            },
            {
                'اسم الصنف في أرصدة المخازن': 'صنف وفر 10',
                'اسم الصنف في الكتالوج': 'Item 10',
                'سعر الجمهور للوحدة': '100.00 ج.م',
                'الخصم المرجح %': '10.0%',
                'خصم هذا المخزن %': '20.0%',
                'أعلى خصم متاح %': '20.0%',
                'المخزن صاحب أعلى خصم': 'مخزن السلام فارم',
                'فرق السعر (الوفر بالجنيه)': '10.00 ج.م'
            }
        ])
        payload = {
            'master_comparison': pd.DataFrame(),
            'per_warehouse_above_average_discounts': {'مخزن السلام فارم': cand_df},
            'allocated_orders': {
                'مخزن السلام فارم': pd.DataFrame([{'م': 1, 'اسم الصنف في أمر الشراء': 'x', 'اسم الصنف في كتالوج المخزن': 'x', 'الكمية المطلوبة': 1, 'نسبة الخصم %': 0.2, 'سعر الجمهور للوحدة': 10, 'سعر التوريد الصافي': 8, 'إجمالي القيمة بسعر الجمهور': 10, 'إجمالي القيمة الصافية': 8, 'ملاحظات التوريد': ''}])
            }
        }
        with tempfile.TemporaryDirectory() as tmp_dir:
            gen = ArabicExcelReportGenerator(payload, output_dir=tmp_dir)
            p = gen.generate(custom_date_str='2026-10-19')
            wb = openpyxl.load_workbook(p)
            ws = [wb[s] for s in wb.sheetnames if 'السلام' in s][0]
            sec_hdr = [
                r for r in range(1, ws.max_row + 1)
                if ws.cell(row=r, column=1).value == 'م' and ws.cell(row=r, column=2).value == 'اسم الصنف في أرصدة المخازن'
            ][0]
            # Row 1 must be صنف وفر 10 (10.00 EGP)
            assert ws.cell(row=sec_hdr + 1, column=2).value == 'صنف وفر 10'
            assert ws.cell(row=sec_hdr + 1, column=9).value == 10.0
            # Row 2 must be صنف وفر 5 (5.00 EGP)
            assert ws.cell(row=sec_hdr + 2, column=2).value == 'صنف وفر 5'
            assert ws.cell(row=sec_hdr + 2, column=9).value == 5.0

    def test_secondary_table_tie_both_styled_soft_green(self):
        """
        Tie Gate: When two warehouses offer the exact same highest discount (e.g. 25%),
        both qualify as the highest catalog discount across all warehouses and are styled
        Soft Green (#E2EFDA) on their respective sheets.
        """
        inv = pd.DataFrame([
            {
                'warehouse_name': 'مخزن السلام فارم',
                'item_name': 'Panadol Extra 24 Tab',
                'normalized_name': 'panadol extra 24 tab',
                'dosage': '', 'unit_size': '',
                'base_price': 100.0, 'discount_rate': 0.25, 'stock_status': True
            },
            {
                'warehouse_name': 'مخزن الفرسان بنها',
                'item_name': 'Panadol Extra 24 Tab',
                'normalized_name': 'panadol extra 24 tab',
                'dosage': '', 'unit_size': '',
                'base_price': 100.0, 'discount_rate': 0.25, 'stock_status': True
            },
        ])
        ref_df = pd.DataFrame([{'clean_name': 'بانادول اكسترا 24 قرص', 'avg_discount_rate': 0.15}])
        alloc = MedicineAllocator(inv, enable_ai=False)
        res = evaluate_above_average_discounts(ref_df, alloc)
        per_wh = res.attrs.get('per_warehouse', {})

        salam_sec = per_wh.get('مخزن السلام فارم', pd.DataFrame())
        forsan_sec = per_wh.get('مخزن الفرسان بنها', pd.DataFrame())

        assert len(salam_sec) == 1
        assert len(forsan_sec) == 1
        assert bool(salam_sec.iloc[0]['_is_winner']) is True
        assert bool(forsan_sec.iloc[0]['_is_winner']) is True

        payload = {
            'master_comparison': pd.DataFrame(),
            'per_warehouse_above_average_discounts': per_wh,
            'allocated_orders': {
                'مخزن السلام فارم': pd.DataFrame([{'م': 1, 'اسم الصنف في أمر الشراء': 'x', 'اسم الصنف في كتالوج المخزن': 'x', 'الكمية المطلوبة': 1, 'نسبة الخصم %': 0.2, 'سعر الجمهور للوحدة': 10, 'سعر التوريد الصافي': 8, 'إجمالي القيمة بسعر الجمهور': 10, 'إجمالي القيمة الصافية': 8, 'ملاحظات التوريد': ''}]),
                'مخزن الفرسان بنها': pd.DataFrame([{'م': 1, 'اسم الصنف في أمر الشراء': 'x', 'اسم الصنف في كتالوج المخزن': 'x', 'الكمية المطلوبة': 1, 'نسبة الخصم %': 0.2, 'سعر الجمهور للوحدة': 10, 'سعر التوريد الصافي': 8, 'إجمالي القيمة بسعر الجمهور': 10, 'إجمالي القيمة الصافية': 8, 'ملاحظات التوريد': ''}]),
            }
        }
        with tempfile.TemporaryDirectory() as tmp_dir:
            gen = ArabicExcelReportGenerator(payload, output_dir=tmp_dir)
            p = gen.generate(custom_date_str='2026-10-19')
            wb = openpyxl.load_workbook(p)
            ws_salam = [wb[s] for s in wb.sheetnames if 'السلام' in s][0]
            ws_forsan = [wb[s] for s in wb.sheetnames if 'الفرسان' in s][0]

            hdr_salam = [r for r in range(1, ws_salam.max_row + 1) if ws_salam.cell(row=r, column=1).value == 'م' and ws_salam.cell(row=r, column=2).value == 'اسم الصنف في أرصدة المخازن'][0]
            hdr_forsan = [r for r in range(1, ws_forsan.max_row + 1) if ws_forsan.cell(row=r, column=1).value == 'م' and ws_forsan.cell(row=r, column=2).value == 'اسم الصنف في أرصدة المخازن'][0]

            # Both must be Soft Green
            assert ws_salam.cell(row=hdr_salam + 1, column=8).fill.start_color.rgb in ['E2EFDA', '00E2EFDA', 'FFE2EFDA']
            assert ws_forsan.cell(row=hdr_forsan + 1, column=8).fill.start_color.rgb in ['E2EFDA', '00E2EFDA', 'FFE2EFDA']
            assert ws_salam.cell(row=hdr_salam + 1, column=8).value == 'مخزن السلام فارم'
            assert ws_forsan.cell(row=hdr_forsan + 1, column=8).value == 'مخزن الفرسان بنها'



