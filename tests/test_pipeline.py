# -*- coding: utf-8 -*-
"""
Test Suite: Pipeline & Controller Bridge (Tiers 1 & 2)
Opaque-box requirement-driven testing covering:
  - Tier 1: File Auto-Categorization, Progress Callback, Pipeline Execution, Ingestion
  - Tier 2: Boundary Value Analysis, Corner Cases, Locks, Corrupt Data, and Error Diagnostics
"""
import os
import sys
import tempfile
import pytest
import pandas as pd
import openpyxl

from src.controller import classify_files, ProcurementEngineBridge
from src.pipeline import (
    run_procurement_pipeline,
    PipelineExecutionError,
    parse_excel_catalog,
)


# ============================================================================
# FIXTURES & TEST DATA HELPERS
# ============================================================================

@pytest.fixture
def sample_catalog_pdfs():
    """Returns list of real sample catalog PDF paths if available."""
    base_dir = os.path.join(os.path.dirname(__file__), "..", "input", "catalogs")
    files = [
        os.path.abspath(os.path.join(base_dir, f))
        for f in os.listdir(base_dir)
        if f.endswith(".pdf")
    ]
    return files


@pytest.fixture
def sample_order_xlsx():
    """Returns path to the primary sample order xlsx."""
    base_dir = os.path.join(os.path.dirname(__file__), "..", "input", "orders")
    target = os.path.abspath(os.path.join(base_dir, "طلبية_اليوم.xlsx"))
    assert os.path.exists(target), f"Sample order not found: {target}"
    return target


@pytest.fixture
def temp_output_dir():
    """Creates an isolated temporary output directory."""
    with tempfile.TemporaryDirectory() as tmpdir:
        yield tmpdir


# ============================================================================
# TIER 1: FEATURE COVERAGE TESTS
# ============================================================================

class TestTier1FileClassification:
    """Tier 1: Comprehensive tests for 3-tier file classification heuristic."""

    def test_classify_pdf_as_catalog(self, temp_output_dir):
        """Tier 1: PDF files are unconditionally classified as catalogs."""
        pdf1 = os.path.join(temp_output_dir, "test_wholesaler.pdf")
        pdf2 = os.path.join(temp_output_dir, "مخزن_الأدوية.pdf")
        with open(pdf1, "wb") as f:
            f.write(b"%PDF-1.4 dummy content")
        with open(pdf2, "wb") as f:
            f.write(b"%PDF-1.4 dummy content")

        result = classify_files([pdf1, pdf2])
        assert len(result["catalogs"]) == 2
        assert pdf1 in result["catalogs"]
        assert pdf2 in result["catalogs"]
        assert len(result["orders"]) == 0
        assert len(result["unknown"]) == 0

    def test_classify_order_by_date_regex(self, temp_output_dir):
        """Tier 1: Spreadsheet files matching 4-digit date regex (e.g. 0910.xlsx) are orders."""
        order1 = os.path.join(temp_output_dir, "0910.xlsx")
        order2 = os.path.join(temp_output_dir, "1231.csv")
        pd.DataFrame([{"الصنف": "Panadol", "الكمية": 5}]).to_excel(order1, index=False)
        pd.DataFrame([{"الصنف": "Brufen", "الكمية": 2}]).to_csv(order2, index=False)

        result = classify_files([order1, order2])
        assert order1 in result["orders"]
        assert order2 in result["orders"]
        assert len(result["catalogs"]) == 0

    def test_classify_order_by_keywords(self, temp_output_dir):
        """Tier 1: Files with order-related keywords in filename are classified as orders."""
        kw_files = [
            "طلبية_اليوم.xlsx",
            "نواقص_الصيدلية.xlsx",
            "purchase_order_september.xlsx",
            "daily_buying_list.csv",
        ]
        paths = []
        for name in kw_files:
            p = os.path.join(temp_output_dir, name)
            if name.endswith(".xlsx"):
                pd.DataFrame([{"الصنف": "Cataflam", "الكمية": 10}]).to_excel(p, index=False)
            else:
                pd.DataFrame([{"الصنف": "Cataflam", "الكمية": 10}]).to_csv(p, index=False)
            paths.append(p)

        result = classify_files(paths)
        for p in paths:
            assert p in result["orders"], f"Expected {p} to be classified as order"
        assert len(result["catalogs"]) == 0

    def test_classify_order_by_content_peek(self, temp_output_dir):
        """Tier 1: Files with ambiguous names (e.g. data_export.xlsx) but order content are classified as orders."""
        ambiguous_path = os.path.join(temp_output_dir, "export_data_778.xlsx")
        df = pd.DataFrame([
            {"أسم العميل": "صيدلية النور", "الكمية المطلوبة": 10, "الصنف": "Concor 5mg"},
            {"أسم العميل": "صيدلية النور", "الكمية المطلوبة": 4, "الصنف": "Nexium 40mg"},
        ])
        df.to_excel(ambiguous_path, index=False)

        result = classify_files([ambiguous_path])
        assert ambiguous_path in result["orders"]
        assert len(result["catalogs"]) == 0

    def test_classify_catalog_negative_keywords(self, temp_output_dir):
        """Tier 1: Files with catalog/wholesaler keywords are classified as catalogs, not orders."""
        cat_files = [
            "عروض_المخازن.xlsx",
            "خصم_سبتمبر_فارما.xlsx",
            "مخزن_السلام_جملة.xlsx",
            "catalog_ex_prices.csv",
        ]
        paths = []
        for name in cat_files:
            p = os.path.join(temp_output_dir, name)
            if name.endswith(".xlsx"):
                pd.DataFrame([{"الصنف": "Panadol", "الخصم": "25%", "السعر": 40}]).to_excel(p, index=False)
            else:
                pd.DataFrame([{"الصنف": "Panadol", "الخصم": "25%", "السعر": 40}]).to_csv(p, index=False)
            paths.append(p)

        result = classify_files(paths)
        for p in paths:
            assert p in result["catalogs"], f"Expected {p} to be catalog"
        assert len(result["orders"]) == 0

    def test_classify_unknown_and_system_files(self, temp_output_dir):
        """Tier 1: Temporary Excel lock files, unsupported extensions, or metadata go to unknown."""
        unsupported = [
            os.path.join(temp_output_dir, "notes.txt"),
            os.path.join(temp_output_dir, "inventory.docx"),
            os.path.join(temp_output_dir, "config.ini"),
            os.path.join(temp_output_dir, "~$daily_order.xlsx"),
            os.path.join(temp_output_dir, "Procurement_Order_Optimized_2026-09-10.xlsx"),
        ]
        for p in unsupported:
            with open(p, "w", encoding="utf-8") as f:
                f.write("placeholder text")

        result = classify_files(unsupported)
        assert len(result["unknown"]) == len(unsupported)
        assert len(result["orders"]) == 0
        assert len(result["catalogs"]) == 0

    def test_classify_mixed_batch(self, temp_output_dir):
        """Tier 1: A mixed batch containing orders, pdf catalogs, excel catalogs, and unknown files."""
        pdf_cat = os.path.join(temp_output_dir, "ibn_sina.pdf")
        with open(pdf_cat, "wb") as f:
            f.write(b"%PDF-1.4 dummy")

        xlsx_cat = os.path.join(temp_output_dir, "عروض_مميز.xlsx")
        pd.DataFrame([{"اسم الصنف": "Augmentin", "نسبة الخصم": 0.25, "سعر الجمهور": 120}]).to_excel(xlsx_cat, index=False)

        xlsx_order = os.path.join(temp_output_dir, "طلبية_اليوم.xlsx")
        pd.DataFrame([{"اسم الصنف": "Augmentin", "الكمية المطلوبة": 5}]).to_excel(xlsx_order, index=False)

        txt_file = os.path.join(temp_output_dir, "readme.txt")
        with open(txt_file, "w") as f:
            f.write("instructions")

        result = classify_files([pdf_cat, xlsx_cat, xlsx_order, txt_file])
        assert pdf_cat in result["catalogs"]
        assert xlsx_cat in result["catalogs"]
        assert xlsx_order in result["orders"]
        assert txt_file in result["unknown"]


class TestTier1PipelineExecution:
    """Tier 1: Core execution pipeline, progress tracking, and Excel output generation."""

    def test_parse_excel_catalog_unit(self, temp_output_dir):
        """Tier 1: parse_excel_catalog extracts items, parses discounts, and normalizes schema."""
        catalog_path = os.path.join(temp_output_dir, "vendor_sample.xlsx")
        raw_data = [
            {"اسم الصنف": "Panadol Extra 24 Tab", "نسبة الخصم": "25%", "سعر الجمهور": 45.0},
            {"اسم الصنف": "Augmentin 1gm 14 Tab", "نسبة الخصم": "0.28", "سعر الجمهور": 120.0},
            {"اسم الصنف": "Cataflam 50mg 20 Tab", "نسبة الخصم": 20.0, "سعر الجمهور": 55.0},
        ]
        pd.DataFrame(raw_data).to_excel(catalog_path, index=False)

        df = parse_excel_catalog(catalog_path, "مخزن التجربة")
        assert not df.empty
        assert len(df) == 3
        expected_cols = [
            "warehouse_name", "item_name", "normalized_name", "dosage",
            "unit_size", "base_price", "discount_rate", "stock_status"
        ]
        for col in expected_cols:
            assert col in df.columns

        # Verify discount normalization
        assert df.loc[df["item_name"].str.contains("Panadol"), "discount_rate"].iloc[0] == pytest.approx(0.25, abs=0.01)
        assert df.loc[df["item_name"].str.contains("Augmentin"), "discount_rate"].iloc[0] == pytest.approx(0.28, abs=0.01)

    def test_pipeline_single_order_single_catalog(self, sample_catalog_pdfs, sample_order_xlsx, temp_output_dir):
        """Tier 1: Single order + single PDF catalog executes successfully and yields valid output files."""
        assert len(sample_catalog_pdfs) > 0, "Requires at least 1 sample catalog PDF"
        cat = sample_catalog_pdfs[0]

        res = run_procurement_pipeline(
            order_files=[sample_order_xlsx],
            catalog_files=[cat],
            output_dir=temp_output_dir,
        )

        assert isinstance(res, dict)
        assert "results" in res
        assert "report_path" in res
        assert "arabic_report_path" in res
        assert "kpis" in res

        assert os.path.isfile(res["report_path"])
        assert os.path.getsize(res["report_path"]) > 1000

        kpis = res["kpis"]
        assert kpis["total_items"] > 0
        assert kpis["matched_items"] + kpis["out_of_stock"] == kpis["total_items"]

    def test_pipeline_multiple_catalogs(self, sample_catalog_pdfs, sample_order_xlsx, temp_output_dir):
        """Tier 1: Multiple catalogs are ingested, consolidated, and competitively allocated."""
        assert len(sample_catalog_pdfs) >= 2, "Requires at least 2 sample catalogs"

        res = run_procurement_pipeline(
            order_files=[sample_order_xlsx],
            catalog_files=sample_catalog_pdfs,
            output_dir=temp_output_dir,
        )

        kpis = res["kpis"]
        assert kpis["total_items"] == 8
        assert kpis["matched_items"] > 0
        assert kpis["avg_discount"] > 0

        # Verify Excel workbook exists and is readable
        wb = openpyxl.load_workbook(res["report_path"])
        assert "المقارنة الشاملة وتوزيع الأصناف" in wb.sheetnames
        assert "نواقص غير متوفرة" in wb.sheetnames

    def test_pipeline_progress_callback_invocations(self, sample_catalog_pdfs, sample_order_xlsx, temp_output_dir):
        """Tier 1: Progress callback is called throughout execution with non-decreasing percentages."""
        history = []

        def track_progress(pct: float, msg: str):
            history.append((pct, msg))

        res = run_procurement_pipeline(
            order_files=[sample_order_xlsx],
            catalog_files=[sample_catalog_pdfs[0]],
            output_dir=temp_output_dir,
            progress_callback=track_progress,
        )

        assert len(history) >= 4, f"Expected at least 4 progress milestones, got {len(history)}"
        pcts = [item[0] for item in history]

        # Verify progress bounds and monotonicity
        assert pcts[0] >= 0.0
        assert pcts[-1] == pytest.approx(1.0, abs=0.01)
        for i in range(1, len(pcts)):
            assert pcts[i] >= pcts[i - 1], f"Progress not monotonic: {pcts[i-1]} -> {pcts[i]}"

        # Verify messages are non-empty strings
        for _, msg in history:
            assert isinstance(msg, str) and len(msg.strip()) > 0

    def test_pipeline_custom_date_string(self, sample_catalog_pdfs, sample_order_xlsx, temp_output_dir):
        """Tier 1: Providing custom_date_str correctly reflects in generated filename."""
        custom_date = "2026-11-20"
        res = run_procurement_pipeline(
            order_files=[sample_order_xlsx],
            catalog_files=[sample_catalog_pdfs[0]],
            output_dir=temp_output_dir,
            custom_date_str=custom_date,
        )
        assert custom_date in os.path.basename(res["report_path"])


# ============================================================================
# TIER 2: BOUNDARY VALUE ANALYSIS & CORNER CASES
# ============================================================================

class TestTier2BoundaryAndCornerCases:
    """Tier 2: Boundary conditions, empty inputs, non-existent paths, locks, and corrupted files."""

    def test_classify_empty_input_list(self):
        """Tier 2: classify_files with empty list returns clean empty structure."""
        result = classify_files([])
        assert result == {"orders": [], "catalogs": [], "unknown": []}

    def test_classify_nonexistent_and_invalid_paths(self):
        """Tier 2: Non-existent files or invalid types are safely relegated to unknown."""
        files = [
            "nonexistent_catalog.pdf",
            "ghost_order.xlsx",
            "",
            None,
        ]
        result = classify_files(files)
        assert len(result["orders"]) == 0
        assert len(result["catalogs"]) == 0
        assert len(result["unknown"]) == len(files)

    def test_classify_case_insensitive_extensions(self, temp_output_dir):
        """Tier 2: Uppercase extensions (.PDF, .XLSX, .CSV) are classified identically."""
        pdf_upper = os.path.join(temp_output_dir, "CATALOG.PDF")
        xlsx_upper = os.path.join(temp_output_dir, "ORDER_0910.XLSX")
        with open(pdf_upper, "wb") as f:
            f.write(b"%PDF-1.4 dummy")
        pd.DataFrame([{"الصنف": "Panadol", "الكمية": 5}]).to_excel(xlsx_upper, engine="openpyxl", index=False)


        result = classify_files([pdf_upper, xlsx_upper])
        assert pdf_upper in result["catalogs"]
        assert xlsx_upper in result["orders"]

    def test_classify_zero_byte_files(self, temp_output_dir):
        """Tier 2: 0-byte spreadsheet files do not crash the content inspection parser."""
        empty_xlsx = os.path.join(temp_output_dir, "empty_file.xlsx")
        empty_csv = os.path.join(temp_output_dir, "empty_file.csv")
        with open(empty_xlsx, "wb") as f:
            pass
        with open(empty_csv, "wb") as f:
            pass

        result = classify_files([empty_xlsx, empty_csv])
        # Since files are 0 bytes and have generic names, should safely default to unknown
        assert empty_xlsx in result["unknown"]
        assert empty_csv in result["unknown"]

    def test_pipeline_empty_order_files_raises_error(self, sample_catalog_pdfs, temp_output_dir):
        """Tier 2: Running pipeline with empty order_files raises PipelineExecutionError or ValueError."""
        with pytest.raises((PipelineExecutionError, ValueError)) as exc_info:
            run_procurement_pipeline(
                order_files=[],
                catalog_files=[sample_catalog_pdfs[0]],
                output_dir=temp_output_dir,
            )
        assert "purchase order" in str(exc_info.value).lower() or "order" in str(exc_info.value).lower()

    def test_pipeline_empty_catalog_files_raises_error(self, sample_order_xlsx, temp_output_dir):
        """Tier 2: Running pipeline with empty catalog_files raises PipelineExecutionError or ValueError."""
        with pytest.raises((PipelineExecutionError, ValueError)) as exc_info:
            run_procurement_pipeline(
                order_files=[sample_order_xlsx],
                catalog_files=[],
                output_dir=temp_output_dir,
            )
        assert "catalog" in str(exc_info.value).lower()

    def test_pipeline_both_empty_raises_error(self, temp_output_dir):
        """Tier 2: Running pipeline with both lists empty raises error."""
        with pytest.raises((PipelineExecutionError, ValueError)):
            run_procurement_pipeline(
                order_files=[],
                catalog_files=[],
                output_dir=temp_output_dir,
            )

    def test_pipeline_nonexistent_order_file(self, sample_catalog_pdfs, temp_output_dir):
        """Tier 2: Specifying an order path that does not exist raises diagnostic error."""
        with pytest.raises(PipelineExecutionError) as exc_info:
            run_procurement_pipeline(
                order_files=["/invalid/path/nonexistent_order.xlsx"],
                catalog_files=[sample_catalog_pdfs[0]],
                output_dir=temp_output_dir,
            )
        assert "order" in str(exc_info.value).lower()

    def test_pipeline_nonexistent_catalog_file(self, sample_order_xlsx, temp_output_dir):
        """Tier 2: Specifying only non-existent catalog files raises diagnostic error."""
        with pytest.raises(PipelineExecutionError) as exc_info:
            run_procurement_pipeline(
                order_files=[sample_order_xlsx],
                catalog_files=["/invalid/path/nonexistent_catalog.pdf"],
                output_dir=temp_output_dir,
            )
        assert "catalog" in str(exc_info.value).lower()

    def test_pipeline_duplicate_items_in_order_aggregated(self, sample_catalog_pdfs, temp_output_dir):
        """Tier 2: Duplicate items in order are summed together into a single allocation row."""
        dup_order_path = os.path.join(temp_output_dir, "order_with_duplicates.xlsx")
        dup_data = [
            {"اسم الصنف": "Panadol Extra 24 Tab", "الكمية المطلوبة": 10},
            {"اسم الصنف": "Panadol Extra 24 Tab", "الكمية المطلوبة": 15},
            {"اسم الصنف": "Brufen 400mg 30 Tab", "الكمية المطلوبة": 4},
        ]
        pd.DataFrame(dup_data).to_excel(dup_order_path, index=False)

        res = run_procurement_pipeline(
            order_files=[dup_order_path],
            catalog_files=[sample_catalog_pdfs[0]],
            output_dir=temp_output_dir,
        )

        kpis = res["kpis"]
        # 2 unique items
        assert kpis["total_items"] == 2

        # Check master comparison quantity for Panadol
        master_df = res["results"]["master_comparison"]
        panadol_row = master_df[master_df["اسم الصنف في أمر الشراء"].str.contains("Panadol", na=False)]
        if not panadol_row.empty:
            assert panadol_row["الكمية"].iloc[0] == 25

    def test_pipeline_all_items_out_of_stock(self, sample_catalog_pdfs, temp_output_dir):
        """Tier 2: Order where 100% of items are not in any catalog -> cleanly routed to out_of_stock."""
        oos_order_path = os.path.join(temp_output_dir, "oos_order.xlsx")
        oos_data = [
            {"اسم الصنف": "NonExistentDrug1 500mg Tab", "الكمية المطلوبة": 3},
            {"اسم الصنف": "UnknownHerbExtract 100ml Syrup", "الكمية المطلوبة": 7},
        ]
        pd.DataFrame(oos_data).to_excel(oos_order_path, index=False)

        res = run_procurement_pipeline(
            order_files=[oos_order_path],
            catalog_files=[sample_catalog_pdfs[0]],
            output_dir=temp_output_dir,
        )

        kpis = res["kpis"]
        assert kpis["total_items"] == 2
        assert kpis["matched_items"] == 0
        assert kpis["out_of_stock"] == 2

        oos_df = res["results"]["out_of_stock"]
        assert len(oos_df) == 2

    def test_pipeline_file_lock_permission_error_handling(self, sample_catalog_pdfs, sample_order_xlsx, temp_output_dir):
        """Tier 2: When the default output file is locked/open in Excel, falls back to timestamped file."""
        import datetime
        today_str = datetime.date.today().strftime("%Y-%m-%d")
        locked_filename = f"Procurement_Order_Optimized_{today_str}.xlsx"
        locked_filepath = os.path.join(temp_output_dir, locked_filename)

        # Simulate Excel holding an open write lock on Windows
        with open(locked_filepath, "w") as lock_file:
            lock_file.write("simulated active lock in excel")
            lock_file.flush()

            # Pipeline should handle PermissionError and write fallback
            res = run_procurement_pipeline(
                order_files=[sample_order_xlsx],
                catalog_files=[sample_catalog_pdfs[0]],
                output_dir=temp_output_dir,
            )

            assert os.path.exists(res["report_path"])
            # Fallback path should differ from locked path or exist with content
            assert os.path.getsize(res["report_path"]) > 1000

    def test_pipeline_cancellation_check(self, sample_catalog_pdfs, sample_order_xlsx, temp_output_dir):
        """Tier 2: cancel_check returning True immediately aborts pipeline with InterruptedError."""
        cancel_flag = [False]

        def should_cancel():
            return cancel_flag[0]

        # Trigger cancellation immediately
        cancel_flag[0] = True
        with pytest.raises((InterruptedError, PipelineExecutionError)):
            run_procurement_pipeline(
                order_files=[sample_order_xlsx],
                catalog_files=[sample_catalog_pdfs[0]],
                output_dir=temp_output_dir,
                cancel_check=should_cancel,
            )

    def test_procurement_engine_bridge_async_lifecycle(self, sample_catalog_pdfs, sample_order_xlsx, temp_output_dir):
        """Tier 2: ProcurementEngineBridge runs in daemon thread and delivers queue messages."""
        bridge = ProcurementEngineBridge()
        assert not bridge.is_running

        bridge.start(
            order_files=[sample_order_xlsx],
            catalog_files=[sample_catalog_pdfs[0]],
            output_dir=temp_output_dir,
        )
        assert bridge.is_running

        bridge.join(timeout=30.0)
        assert not bridge.is_running
        assert bridge.last_result is not None

        # Verify messages arrived in queue
        messages = []
        while True:
            msg = bridge.get_message(block=False)
            if msg is None:
                break
            messages.append(msg)

        msg_types = [m[0] for m in messages]
        assert "PROGRESS" in msg_types
        assert "DONE" in msg_types
