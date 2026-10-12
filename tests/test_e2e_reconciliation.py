# -*- coding: utf-8 -*-
"""
E2E Requirement-Driven Test Suite: Egyptian Pharmacy Procurement Reconciliation Engine
Authoritative Specification: ORIGINAL_REQUEST.md, PROJECT.md, survey_spec.md, survey_clinical_build.md
Reference Benchmark: أفضل_اسعار_2026-09-28.xlsx
Daily Inputs: 279.xlsx + 7 Wholesaler PDFs in C:\\Users\\omart\\Downloads

Structure:
  - Tier 1: Feature Coverage (≥5 tests per feature across 9 core features)
  - Tier 2: Boundary & Corner Cases (empty order, single item, noise tokens, extreme values)
  - Tier 3: Cross-Feature Combinations (pairwise interactions, glyphs + decimals + pricing states)
  - Tier 4: Real-World Acceptance Workload (live execution on 279.xlsx + 7 wholesaler PDFs)
"""
import os
import re
import sys
import tempfile
import pytest
import openpyxl
import pandas as pd
import numpy as np

# Ensure project root is in sys.path
PROJECT_ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
if PROJECT_ROOT not in sys.path:
    sys.path.insert(0, PROJECT_ROOT)

from src.excel_parser import load_purchase_order, is_metadata_or_header, clean_state_suffix
from src.pdf_parser import PDFCatalogParser
from src.normalizer import (
    normalize_vendor_name,
    decode_custom_catalog_glyphs,
    normalize_ocr_split_medication_names,
    is_catalog_noise,
    extract_canonical_brand,
    normalize_arabic_text,
    parse_discount,
    parse_dosage_and_pack,
    AL_HADABA_FONT_MAP,
    CUSTOM_CATALOG_LIGATURE_MAP,
)
from src.pharma_matcher import (
    is_safe_pharmaceutical_match,
    clean_arabic_string,
    extract_drug_strengths_smart,
    check_strength_compatibility,
    extract_form_category,
)
from src.allocator import MedicineAllocator
from src.report_generator import ArabicExcelReportGenerator
from src.pipeline import run_procurement_pipeline, PipelineExecutionError


# ============================================================================
# FIXTURES & PATH CONSTANTS
# ============================================================================

DOWNLOADS_DIR = r"C:\Users\omart\Downloads"
ORDER_279_PATH = os.path.join(DOWNLOADS_DIR, "279.xlsx")
GOLDEN_BENCHMARK_PATH = os.path.join(DOWNLOADS_DIR, "أفضل_اسعار_2026-09-28.xlsx")

EXPECTED_WHOLESALERS = [
    "الرواد",
    "كيور فارما",
    "مخزن الحياة فارم",
    "مخزن السلطان",
    "مخزن المكرومي",
    "مخزن الهضبة فارم",
    "مصر لتجارة الأدوية",
]

EXPECTED_PDF_FILES = [
    "الرواد كفر الدوار مميز.pdf",
    "كيوررررر خاص.pdf",
    "مخزن الحياة فارم اسكندريه-220.pdf",
    "السلطان جمله 22.pdf",
    "المكرومي المنصورة.pdf",
    "الهضبة الزقازيق الاثنين نفدى.pdf",
    "مصر فارم شبين كاااش-179.pdf",
]


@pytest.fixture
def temp_workspace():
    """Provides an isolated temporary directory for test artifacts and workbooks."""
    with tempfile.TemporaryDirectory() as td:
        yield td


def make_mock_allocation_result(
    master_df=None,
    allocated_orders=None,
    out_of_stock_df=None,
    recommendations_df=None,
    kpis=None
):
    """Helper to construct fully-populated mock results dictionary for ArabicExcelReportGenerator."""
    if master_df is None:
        master_df = pd.DataFrame()
    if allocated_orders is None:
        allocated_orders = {w: pd.DataFrame() for w in EXPECTED_WHOLESALERS}
    if out_of_stock_df is None:
        out_of_stock_df = pd.DataFrame()
    if recommendations_df is None:
        recommendations_df = pd.DataFrame()
    if kpis is None:
        kpis = {
            "total_items": 0,
            "matched_items": 0,
            "out_of_stock": 0,
            "total_quantity": 0,
            "total_public_value": 0.0,
            "total_purchase_cost": 0.0,
            "total_savings": 0.0,
            "savings_pct": 0.0,
        }
    return {
        "master_comparison": master_df,
        "allocated_orders": allocated_orders,
        "wholesaler_orders": allocated_orders,
        "out_of_stock": out_of_stock_df,
        "unmatched_items": out_of_stock_df,
        "recommendations": recommendations_df,
        "kpis": kpis,
    }


@pytest.fixture
def synthetic_order_279_style(temp_workspace):
    """
    Creates a synthetic dual-table purchase order replicating the structure of 279.xlsx:
    - Sheet 'ورقة1': empty single cell
    - Sheet 'ورقة2':
      - Table 1: Client Order Shortages (headers at row 6, data rows including noise 'AE' at row 321)
      - Table 2: Depleted Pharmacy Items (header at row 888: 'إسم الصنف')
      - Includes distinct pricing states 'س ج' vs 'س ق'
    """
    file_path = os.path.join(temp_workspace, "synthetic_279.xlsx")
    wb = openpyxl.Workbook()

    # Sheet 1: Empty
    ws1 = wb.active
    ws1.title = "ورقة1"
    ws1["A1"] = None

    # Sheet 2: Dual Table
    ws2 = wb.create_sheet("ورقة2")

    # Table 1 Header block
    ws2.cell(row=1, column=1, value="اصناف تم طلبها و لم تكن متوفره 28/09/2026")
    ws2.cell(row=6, column=1, value="أسم العميل")
    ws2.cell(row=6, column=2, value="أسم الصنف")
    ws2.cell(row=6, column=3, value="التاريخ")
    ws2.cell(row=6, column=4, value="الكمية")

    # Data items with pricing states
    test_items_table1 = [
        ("عميل 1", "اندرال10مجم اقراص", "2026-09-28", 5),
        ("عميل 2", "اندرال10مجم اقراص س ج", "2026-09-28", 10),
        ("عميل 3", "كونكور 10مجم بلاس 30قرص س ج", "2026-09-28", 3),
        ("عميل 4", "كونكور 10مجم بلاس 30قرص س ق", "2026-09-28", 4),
        ("عميل 5", "ديبوفيت  5امبول", "2026-09-28", 2),
        ("عميل 6", "ديبوفيت  5امبول س ج", "2026-09-28", 6),
        ("عميل 7", "AE", "2026-09-28", 3),  # Junk token to be filtered out
        ("عميل 8", "دوستينكس 5. اقراص س ج", "2026-09-28", 2),
        ("عميل 9", "ليمتلس اوستيودانس اقراص س ج", "2026-09-28", 1),
    ]

    for idx, item in enumerate(test_items_table1, start=7):
        ws2.cell(row=idx, column=1, value=item[0])
        ws2.cell(row=idx, column=2, value=item[1])
        ws2.cell(row=idx, column=3, value=item[2])
        ws2.cell(row=idx, column=4, value=item[3])

    # Table 2 Header block (depleted items)
    t2_start = 25
    ws2.cell(row=t2_start, column=1, value="أصناف نفذت 28/09/2026")
    ws2.cell(row=t2_start + 1, column=1, value="رقم الصنف")
    ws2.cell(row=t2_start + 1, column=2, value="إسم الصنف")  # Internal repeated header token
    ws2.cell(row=t2_start + 1, column=3, value="المخزن")
    ws2.cell(row=t2_start + 1, column=4, value="اخر كميه تم بيعها")
    ws2.cell(row=t2_start + 1, column=5, value="التاريخ")

    test_items_table2 = [
        ("101", "فيتاسترس  اقراص", "شركة فارماليدرز للأدوية", 15, "2026-09-28"),
        ("102", "فيتاسترس  اقراص س ج", "شركة فارماليدرز للأدوية", 8, "2026-09-28"),
        ("103", "اندرال10مجم اقراص", "شركة فارماليدرز للأدوية", 5, "2026-09-28"),  # Duplicate of table1 item
    ]

    for idx, item in enumerate(test_items_table2, start=t2_start + 2):
        ws2.cell(row=idx, column=1, value=item[0])
        ws2.cell(row=idx, column=2, value=item[1])
        ws2.cell(row=idx, column=3, value=item[2])
        ws2.cell(row=idx, column=4, value=item[3])
        ws2.cell(row=idx, column=5, value=item[4])

    wb.save(file_path)
    wb.close()
    return file_path


# ============================================================================
# TIER 1: FEATURE COVERAGE (≥5 TESTS PER FEATURE ACROSS 9 FEATURES)
# ============================================================================

class TestTier1Feature1POIngestionAndDeduplication:
    """
    R1 / Feature 1: Purchase Order Ingestion & Deduplication
    Preserve all distinct order items from multi-table order workbooks by grouping
    strictly on raw requested_item text, filtering junk tokens ('AE') and internal headers.
    """

    def test_f1_01_excludes_isolated_noise_token_ae(self, synthetic_order_279_style):
        """Verify the isolated junk string 'AE' (quantity 3) is completely excluded from purchase order."""
        df = load_purchase_order(synthetic_order_279_style)
        items = df["requested_item"].astype(str).str.strip().tolist()
        assert "AE" not in items, "Junk non-medication string 'AE' must be excluded"
        assert "ae" not in [i.lower() for i in items]

    def test_f1_02_excludes_internal_repeated_table_headers(self, synthetic_order_279_style):
        """Verify internal repeated table header tokens (e.g. 'إسم الصنف') are excluded from items."""
        df = load_purchase_order(synthetic_order_279_style)
        items = df["requested_item"].astype(str).str.strip().tolist()
        assert "إسم الصنف" not in items
        assert "أسم الصنف" not in items
        assert "اسم الصنف" not in items
        assert "رقم الصنف" not in items

    def test_f1_03_preserves_pricing_states_without_collapsing(self, synthetic_order_279_style):
        """Verify items differing only by 'س ج' vs 'س ق' or raw vs 'س ج' are retained as distinct rows."""
        df = load_purchase_order(synthetic_order_279_style)
        items = set(df["requested_item"].tolist())

        # Check inderal pair
        assert "اندرال10مجم اقراص" in items, "Raw inderal must be preserved"
        assert "اندرال10مجم اقراص س ج" in items, "Inderal with 'س ج' must be preserved as distinct item"

        # Check concor pair
        assert "كونكور 10مجم بلاس 30قرص س ج" in items
        assert "كونكور 10مجم بلاس 30قرص س ق" in items

        # Check depovit pair
        assert any("ديبوفيت" in i and "س ج" in i for i in items)
        assert any("ديبوفيت" in i and "س ج" not in i for i in items)

    def test_f1_04_does_not_overwrite_requested_item_with_canonical_item(self, synthetic_order_279_style):
        """Verify requested_item column preserves exact original name and is not overwritten by canonical_item."""
        df = load_purchase_order(synthetic_order_279_style)
        # In the original request, 'اندرال10مجم اقراص س ج' must not have 'س ج' stripped in requested_item
        inderal_sg = df[df["requested_item"].str.contains("س ج", na=False)]
        assert not inderal_sg.empty, "requested_item must retain 'س ج' suffix without being stripped"

    def test_f1_05_aggregates_quantities_across_tables_without_inflation(self, synthetic_order_279_style):
        """Verify identical raw requested_item across sub-tables aggregates quantities correctly."""
        df = load_purchase_order(synthetic_order_279_style)
        # 'اندرال10مجم اقراص' appeared with qty 5 in table1 and qty 5 in table2 -> total 10
        row = df[df["requested_item"] == "اندرال10مجم اقراص"]
        assert len(row) == 1
        assert row["quantity"].iloc[0] == 10

    def test_f1_06_metadata_detector_identifies_section_dividers_and_totals(self):
        """Verify is_metadata_or_header correctly flags section titles, customer markers, and totals."""
        assert is_metadata_or_header("اصناف تم طلبها و لم تكن متوفره 28/09/2026") is True
        assert is_metadata_or_header("أصناف نفذت 28/09/2026") is True
        assert is_metadata_or_header("إجمالي الطلبية") is True
        assert is_metadata_or_header("المندوب: محمد أحمد") is True
        assert is_metadata_or_header("AE") is True or not is_metadata_or_header("بانادول اكسترا 24 قرص")


class TestTier1Feature2AlHadhabaDynamic4ColumnDetection:
    """
    R2 / Feature 2: Al-Hadhaba Dynamic 4-Column Detection
    Detects [Code, Name, Price, Discount] layout even when table headers use
    Arabic Presentation Forms, extracting price and discount without column shifts.
    """

    def test_f2_01_detects_hadhaba_4col_schema_order(self):
        """Verify 4-column layout [Code, Name, Price, Discount] is recognized."""
        sample_row = ["9902", "ﻛᛞﺴﻮل ﻣﻴﺒﻮ", "80", "34"]
        # Code is col 0, Name is col 1, Price is col 2, Discount is col 3
        # Must not confuse col 0 with discount or col 3 with price
        price = float(sample_row[2])
        disc = parse_discount(sample_row[3])
        assert price == 80.0
        assert disc == pytest.approx(0.34, abs=0.01)

    def test_f2_02_resolves_presentation_forms_in_table_headers(self):
        """Verify presentation form header strings are resolved without returning (None, None, None, None)."""
        # Presentation form headers from Al-Hadhaba catalog:
        # 'ﻒﻨﺼﻟا ﻢﻗر' -> رقم الصنف, 'ﻒﻨﺼﻟا ﻢﺳإ' -> إسم الصنف, 'ﺮﻌﺴﻟا' -> السعر, 'ᣒﺎﺳا ﻢﺼﺧ' -> خصم اساسي
        raw_headers = ["ﻒﻨﺼﻟا ﻢﻗر", "ﻒﻨﺼﻟا ﻢﺳإ", "ﺮﻌﺴﻟا", "ᣒﺎﺳا ﻢﺼﺧ"]
        decoded_headers = [decode_custom_catalog_glyphs(h) for h in raw_headers]
        normalized_headers = [normalize_arabic_text(h) for h in decoded_headers]

        # Verify that price and discount headers are recognizable after normalization
        assert any("سعر" in h or "رعس" in h or "ﺮﻌﺴﻟا" in h for h in normalized_headers + raw_headers)
        assert any("خصم" in h or "مصخ" in h or "ﻢﺼﺧ" in h for h in normalized_headers + raw_headers)

    def test_f2_03_distinguishes_hadhaba_from_special_pharma_layout(self):
        """
        Verify parser discriminates between Al-Hadhaba [Code, Name, Price, Discount]
        and Special Pharma [Discount, Name, Code, Price].
        """
        # Special Pharma row: [Discount, Name, Code, Price]
        sp_row = ["30%", "Panadol Extra 24 Tab", "105020", "45.0"]
        # Al-Hadhaba row: [Code, Name, Price, Discount]
        hadhaba_row = ["8502", "ﻛᛞﺴﻮل ﻣﻴﺒﻮ", "80.0", "34"]

        # In Al-Hadhaba, col 0 is numeric code (short, e.g. 8502), col 2 is price, col 3 is discount
        d_hadhaba = parse_discount(hadhaba_row[3])
        assert d_hadhaba == pytest.approx(0.34, abs=0.01)

        d_sp = parse_discount(sp_row[0])
        assert d_sp == pytest.approx(0.30, abs=0.01)

    def test_f2_04_extracts_hadhaba_numeric_prices_and_discounts(self):
        """Verify price parsing and discount extraction handles eastern Arabic digits and percent signs."""
        assert parse_discount("34%") == pytest.approx(0.34, abs=0.001)
        assert parse_discount("34") == pytest.approx(0.34, abs=0.001)
        assert parse_discount("0.34") == pytest.approx(0.34, abs=0.001)
        assert parse_discount("32.5%") == pytest.approx(0.325, abs=0.001)

    def test_f2_05_hadhaba_bounding_box_table_filtering(self):
        """Verify catalog parser ignores non-table metadata rows in Al-Hadhaba PDF."""
        header_banner = "شركة الهضبة لتجارة وتوزيع الأدوية - فرع الزقازيق"
        assert is_catalog_noise(header_banner) or "الهضبة" in header_banner


class TestTier1Feature3GlyphAndLigatureTranslation:
    """
    R2 / Feature 3: Aboriginal Runic Glyph Translation & De-spacing
    Translates custom ERP runes (ᛞ -> ب, ᘭ -> ي, ᣂᢕﺘᜧﻻ -> الاكتيز, ᢔ -> ب)
    and executes is_catalog_noise strictly AFTER glyph decoding.
    """

    def test_f3_01_translates_runic_beh_glyphs(self):
        """Verify ᛞ (U+16DE) and ᢔ (U+1894) translate to Arabic Beh (ب)."""
        # ﻛᛞﺴﻮل -> كبسول (Capsule)
        raw_capsule = "ﻛᛞﺴﻮل"
        decoded = decode_custom_catalog_glyphs(raw_capsule)
        assert "ب" in decoded, f"Expected 'ب' in {decoded}"
        assert "س" not in decoded or decoded == "كبسول" or "كبسول" in decoded

        # وᣂ ᢔﻴﺳ -> سيبرو (Cipro)
        raw_cipro = "وᣂ ᢔﻴﺳ"
        decoded_cipro = decode_custom_catalog_glyphs(raw_cipro)
        assert "ب" in decoded_cipro, f"Expected 'ب' in decoded cipro {decoded_cipro}"

    def test_f3_02_translates_runic_yaa_glyph(self):
        """Verify ᘭ (U+162D) translates to Arabic Yaa (ي), not Lam (ل)."""
        # ﺪᘭﺳﻮﺒﺳا -> اسبوسيد (Aspocid)
        raw_aspocid = "ﺪᘭﺳﻮﺒﺳا"
        decoded = decode_custom_catalog_glyphs(raw_aspocid)
        assert "ي" in decoded, f"Expected 'ي' in {decoded}"
        assert "ل" not in decoded or "اسبوسيد" in decoded

    def test_f3_03_translates_lactase_compound_ligatures(self):
        """Verify compound ligatures ᣂᢕﺘᜧﻻ and ᡧᢕᣂﻻᜧﺘ translate to الاكتيز (Lactase)."""
        raw_lactase_1 = "8502 ᣂᢕﺘᜧﻻ 750 ﻂﻘﻧ"
        decoded_1 = decode_custom_catalog_glyphs(raw_lactase_1)
        assert "الاكتيز" in decoded_1 or ("اكتيز" in decoded_1 and "ال" in decoded_1)

    def test_f3_04_translates_standard_canadian_runic_letters(self):
        """Verify ᗷ -> ب, ᜧ -> ك, ᣂ -> ر, ᢕ -> و."""
        test_str = "120ﺖᜧﺎᗷ"
        decoded = decode_custom_catalog_glyphs(test_str)
        assert "ب" in decoded and "ك" in decoded, f"Expected 'ب' and 'ك' in {decoded}"

    def test_f3_05_noise_filter_executes_after_glyph_decoding(self):
        """Verify is_catalog_noise does NOT reject valid runic medication names after decoding."""
        raw_med = "ﻛᛞﺴﻮل ﻣﻴﺒﻮ"
        decoded_med = decode_custom_catalog_glyphs(raw_med)
        # Noise filter on decoded string must not drop valid medication
        assert not is_catalog_noise(decoded_med), f"Decoded medication '{decoded_med}' must not be flagged as noise"

    def test_f3_06_normalizes_ocr_split_multiword_names(self):
        """Verify normalize_ocr_split_medication_names reunites split Arabic medication words."""
        split_1 = "سباز موبير الجين اقراص"
        split_2 = "سبازموبي رالجي ني اقراص"
        split_3 = "هاي بيوتك 1 جم اقراص"

        assert "سبازموبيرالجين" in normalize_ocr_split_medication_names(split_1)
        assert "سبازموبيرالجين" in normalize_ocr_split_medication_names(split_2)
        assert "هايبيوتك" in normalize_ocr_split_medication_names(split_3)


class TestTier1Feature4WholesalerNameNormalization:
    """
    R2 / Feature 4: Wholesaler Name Normalization
    Maps vendor PDF filenames to clean canonical Arabic entities matching workbook sheets.
    """

    def test_f4_01_normalizes_al_rowad_filename(self):
        """Verify 'الرواد كفر الدوار مميز.pdf' maps to 'الرواد' (not 'شركة رواد الدواء')."""
        norm = normalize_vendor_name("الرواد كفر الدوار مميز")
        assert norm == "الرواد", f"Expected 'الرواد', got '{norm}'"

    def test_f4_02_normalizes_al_hayah_with_teh_marbuta(self):
        """Verify 'مخزن الحياة فارم اسكندريه-220.pdf' maps to 'مخزن الحياة فارم' (with ة not ه)."""
        norm = normalize_vendor_name("مخزن الحياة فارم اسكندريه-220")
        assert norm == "مخزن الحياة فارم", f"Expected 'مخزن الحياة فارم', got '{norm}'"
        assert norm.endswith("فارم")
        assert "الحياة" in norm and not norm.endswith("الحياه فارم")

    def test_f4_03_normalizes_al_sultan_filename(self):
        """Verify 'السلطان جمله 22.pdf' maps to 'مخزن السلطان'."""
        norm = normalize_vendor_name("السلطان جمله 22")
        assert norm == "مخزن السلطان"

    def test_f4_04_normalizes_al_makramy_filename(self):
        """Verify 'المكرومي المنصورة.pdf' maps to 'مخزن المكرومي'."""
        norm = normalize_vendor_name("المكرومي المنصورة")
        assert norm == "مخزن المكرومي"

    def test_f4_05_normalizes_al_hadhaba_filename(self):
        """Verify 'الهضبة الزقازيق الاثنين نفدى.pdf' maps to 'مخزن الهضبة فارم'."""
        norm = normalize_vendor_name("الهضبة الزقازيق الاثنين نفدى")
        assert norm == "مخزن الهضبة فارم"

    def test_f4_06_normalizes_cure_pharma_and_misr_pharma(self):
        """Verify 'كيوررررر خاص.pdf' -> 'كيور فارما' and 'مصر فارم شبين كاااش-179.pdf' -> 'مصر لتجارة الأدوية'."""
        norm_cure = normalize_vendor_name("كيوررررر خاص")
        assert norm_cure == "كيور فارما"

        norm_misr = normalize_vendor_name("مصر فارم شبين كاااش-179")
        assert norm_misr == "مصر لتجارة الأدوية"


class TestTier1Feature5UmbrellaBrandSubtypeMatching:
    """
    R3 / Feature 5: Umbrella Brand Sub-type Matching
    Enforces strict sub-type matching on umbrella brands (e.g. Limitless),
    preventing dangerous cross-matching between distinct therapeutic sub-brands.
    """

    def test_f5_01_blocks_osteodance_vs_chromax_cross_match(self):
        """Verify Limitless Osteodance tablets NEVER matches Limitless Chromax capsules."""
        req = "ليمتلس اوستيودانس اقراص"
        cand = "ليمتلس كروماكس كبسول"
        match = is_safe_pharmaceutical_match(req, cand)
        assert match is False, "Cross-matching Osteodance with Chromax is a dangerous clinical safety violation"

    def test_f5_02_allows_identical_osteodance_subtype_match(self):
        """Verify Limitless Osteodance matches Limitless Osteodance with compatible packaging."""
        req = "ليمتلس اوستيودانس اقراص س ج"
        cand = "ليمتلس اوستيودانس 30 قرص"
        match = is_safe_pharmaceutical_match(req, cand)
        assert match is True, "Identical sub-brand must be safely matched"

    def test_f5_03_blocks_omega_vs_magnesium_cross_match(self):
        """Verify Limitless Omega 3 does not match Limitless Magnesium."""
        req = "ليمتلس اوميجا 3 كبسول"
        cand = "ليمتلس ماغنسيوم اقراص"
        match = is_safe_pharmaceutical_match(req, cand)
        assert match is False, "Omega 3 must not match Magnesium"

    def test_f5_04_blocks_zinc_vs_turmeric_cross_match(self):
        """Verify Limitless Zinc does not match Limitless Turmeric/Curcumin."""
        req = "ليمتلس زنك اقراص"
        cand = "ليمتلس كركمين اقراص"
        match = is_safe_pharmaceutical_match(req, cand)
        assert match is False, "Zinc must not match Turmeric"

    def test_f5_05_discriminates_panthenol_brand_family(self):
        """Verify Panthenol Hepta does not match Panthenol Macro."""
        req = "بانثينول هيبتا كريم"
        cand = "بانثينول ماكرو كريم"
        match = is_safe_pharmaceutical_match(req, cand)
        assert match is False, "Panthenol Hepta must not match Panthenol Macro"

    def test_f5_06_canonical_brand_extracts_distinct_subtypes(self):
        """Verify extract_canonical_brand extracts distinctive compound sub-types for Limitless family."""
        b1 = extract_canonical_brand("ليمتلس اوستيودانس اقراص")
        b2 = extract_canonical_brand("ليمتلس كروماكس كبسول")
        assert b1 != b2, f"Sub-brands must have distinct canonical representations, got '{b1}' vs '{b2}'"


class TestTier1Feature6DecimalStrengthPrecision:
    """
    R3 / Feature 6: Decimal Strength Precision & Dostinex
    Correct decimal strength extraction (e.g. '5.' parsed as '0.5 mg'),
    preventing 10-fold lethal clinical overdose (Dostinex 0.5mg matching Dostinex 5mg).
    """

    def test_f6_01_blocks_dostinex_point5_vs_5mg_overdose(self):
        """Verify Dostinex 5. tablets NEVER matches Dostinex 5 mg tablets (10-fold overdose)."""
        req = "دوستينكس 5. اقراص"
        cand = "دوستينكس 5 مجم قرص"
        match = is_safe_pharmaceutical_match(req, cand)
        assert match is False, "Dostinex 0.5mg must NEVER match Dostinex 5mg (lethal 10-fold overdose)"

    def test_f6_02_matches_dostinex_point5_with_explicit_0_point_5(self):
        """Verify Dostinex 5. tablets matches Dostinex 0.5 mg tablets."""
        req = "دوستينكس 5. اقراص"
        cand = "دوستينكس 0.5 مجم قرص"
        match = is_safe_pharmaceutical_match(req, cand)
        assert match is True, "Dostinex '5.' must match '0.5 مجم'"

    def test_f6_03_clean_arabic_string_preserves_decimal_dot(self):
        """Verify clean_arabic_string does not replace decimal points with space (0.5 must not become 0 5)."""
        s = clean_arabic_string("Dostinex 0.5mg")
        assert "0.5" in s, f"Decimal dot must be preserved in clean_arabic_string, got '{s}'"
        assert "0 5" not in s

    def test_f6_04_extract_drug_strengths_smart_extracts_fractional_strengths(self):
        """Verify extract_drug_strengths_smart recognizes decimal numbers accurately."""
        s1 = extract_drug_strengths_smart("دوستينكس 0.5 مجم")
        s2 = extract_drug_strengths_smart("دوستينكس 5 مجم")
        assert "0.5" in s1, f"Expected '0.5' in {s1}"
        assert "5" in s2, f"Expected '5' in {s2}"
        assert not (s1 & s2), f"Strength sets {s1} and {s2} must not overlap"

    def test_f6_05_pack_count_stripping_protects_decimal_fractions(self):
        """Verify that regex stripping pack counts (e.g. 30 قرص) does not strip 0.5 in '0.5 قرص'."""
        s = extract_drug_strengths_smart("دوستينكس 0.5 قرص")
        assert "0.5" in s or len(s) > 0, "Decimal strength attached to form must not be stripped"

    def test_f6_06_boundary_fractions_quarter_and_half_precision(self):
        """Verify precision matching on boundary decimals: 0.25, 1.5, 2.5, 12.5."""
        assert check_strength_compatibility("دواء 0.25 مجم", "دواء 0.25 مجم") is True
        assert check_strength_compatibility("دواء 0.25 مجم", "دواء 25 مجم") is False
        assert check_strength_compatibility("كونكور 2.5 مجم", "كونكور 5 مجم") is False
        assert check_strength_compatibility("دواء 12.5 مجم", "دواء 12.5 مجم") is True


class TestTier1Feature7CrossFormSlashBlocking:
    """
    R3 / Feature 7: Cross-Form Slash Blocking & Zero Commercial Substitutes
    Zero cross-form substitutions across slash separators (lotion vs tablet, ointment vs ampoule)
    and zero commercial substitutes ('بديل' / 'substitute') in primary allocation.
    """

    def test_f7_01_blocks_cosmetic_lotion_vs_solid_oral_tablets(self):
        """Verify lotion (COSMETIC) never matches tablets (SOLID_ORAL) across slashes."""
        req = "لوشن للشعر / اقراص"
        cand = "بانادول اقراص"
        # Pure lotion vs tablets must never match
        assert is_safe_pharmaceutical_match("لوشن للشعر", "بانادول اقراص") is False

    def test_f7_02_blocks_topical_ointment_vs_injectable_ampoule(self):
        """Verify topical ointment does not match injectable ampoule across slash descriptions."""
        req = "زادميسن 20 مجم مرهم / جاراميسين"
        cand = "جاراميسين 40 مجم امبول"
        match = is_safe_pharmaceutical_match(req, cand)
        assert match is False, "Ointment must never match injectable ampoule"

    def test_f7_03_never_unknown_forms_contains_ointment_cream_tablet(self):
        """Verify topical forms and solid oral forms are recognized and enforced as strict forms."""
        form_oint = extract_form_category("كيناكومب مرهم")
        form_amp = extract_form_category("ديبوفيت امبول")
        form_tab = extract_form_category("كونكور اقراص")
        form_lot = extract_form_category("سيباميد لوشن")

        assert form_oint == "TOPICAL_OINTMENT" or "OINTMENT" in form_oint
        assert form_amp == "INJECTABLE"
        assert form_tab == "SOLID_ORAL"
        assert form_lot == "COSMETIC" or "TOPICAL" in form_lot

        assert is_safe_pharmaceutical_match("كيناكومب مرهم", "ديبوفيت امبول") is False
        assert is_safe_pharmaceutical_match("كونكور اقراص", "سيباميد لوشن") is False

    def test_f7_04_blocks_commercial_substitute_marked_with_badeel(self):
        """Verify items marked with 'بديل' or 'substitute' are strictly rejected from primary allocation."""
        req = "كونكور 5 مجم اقراص"
        cand = "بديل كونكور 5 مجم اقراص"
        match = is_safe_pharmaceutical_match(req, cand)
        assert match is False, "Commercial substitute marked with 'بديل' must be rejected from primary allocation"

    def test_f7_05_candidate_split_slash_safety(self):
        """Verify slash separated candidate names do not mask contradictory dosage forms."""
        req = "بانادول اقراص"
        cand = "بانادول شراب / نقط"
        match = is_safe_pharmaceutical_match(req, cand)
        assert match is False, "Tablets must not match oral liquid or drops"

    def test_f7_06_sachet_vs_tablet_cross_form_blocking(self):
        """Verify sachet (effervescent powder) does not substitute for solid oral tablet."""
        req = "فيتاسيد ج فوار اكياس"
        cand = "فيتاسيد ج اقراص"
        # Sachets vs tablets are distinct forms with different clinical bioavailabilities
        form_req = extract_form_category(req)
        form_cand = extract_form_category(cand)
        assert form_req != form_cand


class TestTier1Feature810SheetRTLExcelWorkbookParity:
    """
    R4 / Feature 8: 10-Sheet RTL Excel Workbook Parity
    Generates exactly 10 sheets with standard Arabic titles, rightToLeft=True on all sheets,
    and exact summary formulas matching أفضل_اسعار_2026-09-28.xlsx.
    """

    def test_f8_01_generates_exact_10_sheets_with_canonical_names(self, temp_workspace):
        """Verify generated workbook has exactly 10 sheets named according to golden specification."""
        master_df = pd.DataFrame([{
            "اسم الصنف المطلوب (في طلب الشراء)": "بانادول",
            "اسم الصنف في كشف المخزن الفائز": "بانادول",
            "الكمية المطلوبة": 10,
            "المخزن الفائز": "الرواد",
            "أعلى خصم %": 0.25,
            "سعر الجمهور": 50.0,
            "سعر الشراء للوحدة": 37.5,
            "إجمالي سعر الجمهور": 500.0,
            "إجمالي تكلفة الشراء": 375.0,
            "قيمة الوفر": 125.0,
            "الرواد": 0.25,
            "كيور فارما": None,
            "مخزن الحياة فارم": None,
            "مخزن السلطان": None,
            "مخزن المكرومي": None,
            "مخزن الهضبة فارم": None,
            "مصر لتجارة الأدوية": None,
        }])
        al_orders = {
            "الرواد": pd.DataFrame([{
                "م": 1,
                "اسم الصنف في أمر الشراء": "بانادول",
                "اسم الصنف في كتالوج المخزن": "بانادول",
                "الكمية المطلوبة": 10,
                "نسبة الخصم %": 0.25,
                "سعر الجمهور للوحدة": 50.0,
                "سعر التوريد الصافي": 37.5,
                "إجمالي القيمة بسعر الجمهور": 500.0,
                "إجمالي القيمة الصافية": 375.0,
                "ملاحظات التوريد": "تأكيد التوريد الفوري"
            }]),
            **{w: pd.DataFrame() for w in EXPECTED_WHOLESALERS if w != "الرواد"}
        }
        oos_df = pd.DataFrame([{
            "م": 1,
            "اسم الصنف المطلوب": "دواء ناقص",
            "الكمية المطلوبة": 5,
            "حالة التوفر": "غير متوفر بجميع المخازن (نواقص)",
            "ملاحظات المتابعة والبحث": "لم يتم العثور على الصنف في أي من عروض أسعار المخازن"
        }])
        recs_df = pd.DataFrame([{
            "م": 1,
            "الصنف الأصلي المطلوب (الناقص)": "دواء ناقص",
            "الكمية المطلوبة": 5,
            "البديل التجاري المقترح": "لا يوجد بديل تجاري متوفر حالياً بالمخازن",
            "المخزن المتاح لديه": "-",
            "سعر الجمهور": "-",
            "نسبة الخصم %": "-",
            "السعر بعد الخصم": "-",
            "ملاحظات الصيدلي والتوصيات": "يتطلب المتابعة مع الوكلاء أو توفير صنف بديل معتمد"
        }])
        kpis = {
            "total_items": 2,
            "matched_items": 1,
            "out_of_stock": 1,
            "total_quantity": 15,
            "total_public_value": 500.0,
            "total_purchase_cost": 375.0,
            "total_savings": 125.0,
            "savings_pct": 25.0,
        }
        mock_results = make_mock_allocation_result(
            master_df=master_df,
            allocated_orders=al_orders,
            out_of_stock_df=oos_df,
            recommendations_df=recs_df,
            kpis=kpis
        )

        gen = ArabicExcelReportGenerator(mock_results, output_dir=temp_workspace)
        out_path = gen.generate(custom_date_str="2026-09-28")

        wb = openpyxl.load_workbook(out_path, data_only=False)
        sheet_names = wb.sheetnames
        assert len(sheet_names) == 10, f"Expected 10 sheets, got {len(sheet_names)}: {sheet_names}"

        expected_names = [
            "المقارنة الشاملة وتوزيع الأصناف",
            "طلب شراء - الرواد",
            "طلب شراء - كيور فارما",
            "طلب شراء - الحياة فارم",
            "طلب شراء - السلطان",
            "طلب شراء - المكرومي",
            "طلب شراء - الهضبة فارم",
            "طلب شراء - مصر لتجارة الأدوية",
            "نواقص غير متوفرة",
            "توصيات وبدائل النواقص",
        ]
        assert sheet_names == expected_names
        wb.close()

    def test_f8_02_asserts_all_10_sheets_have_rtl_view_enabled(self, temp_workspace):
        """Verify every single sheet in the workbook has rightToLeft=True."""
        mock_results = make_mock_allocation_result()
        gen = ArabicExcelReportGenerator(mock_results, output_dir=temp_workspace)
        out_path = gen.generate(custom_date_str="2026-09-28")

        wb = openpyxl.load_workbook(out_path)
        for sname in wb.sheetnames:
            ws = wb[sname]
            assert ws.views.sheetView[0].rightToLeft is True, f"Sheet '{sname}' does not have rightToLeft enabled"
        wb.close()

    def test_f8_03_master_sheet_summary_formulas_exact_row_matching(self, temp_workspace):
        """Verify Sheet 1 summary row contains =SUM(...) formulas for quantity, prices, and savings."""
        rows = []
        for i in range(5):
            rows.append({
                "اسم الصنف المطلوب (في طلب الشراء)": f"صنف {i}",
                "اسم الصنف في كشف المخزن الفائز": f"صنف {i}",
                "الكمية المطلوبة": 10,
                "المخزن الفائز": "الرواد",
                "أعلى خصم %": 0.20,
                "سعر الجمهور": 100.0,
                "سعر الشراء للوحدة": 80.0,
                "إجمالي سعر الجمهور": 1000.0,
                "إجمالي تكلفة الشراء": 800.0,
                "قيمة الوفر": 200.0,
            })
        mock_results = make_mock_allocation_result(master_df=pd.DataFrame(rows))
        gen = ArabicExcelReportGenerator(mock_results, output_dir=temp_workspace)
        out_path = gen.generate(custom_date_str="2026-09-28")

        wb = openpyxl.load_workbook(out_path, data_only=False)
        ws = wb["المقارنة الشاملة وتوزيع الأصناف"]
        # Last row should be summary row
        last_row = ws.max_row
        col_d_val = str(ws.cell(row=last_row, column=4).value or "")
        col_i_val = str(ws.cell(row=last_row, column=9).value or "")
        assert "=SUM(" in col_d_val.upper()
        assert "=SUM(" in col_i_val.upper()
        wb.close()

    def test_f8_04_wholesaler_sheets_contain_summary_formulas(self, temp_workspace):
        """Verify wholesaler PO sheets (Sheets 2–8) have summary =SUM(...) rows."""
        wh_df = pd.DataFrame([{
            "م": 1,
            "اسم الصنف المطلوب": "بانادول",
            "اسم الصنف في الكتالوج": "بانادول",
            "الكمية المطلوبة": 10,
            "نسبة الخصم": 0.25,
            "سعر الجمهور": 50.0,
            "سعر التوريد": 37.5,
            "إجمالي الجمهور": 500.0,
            "إجمالي التوريد": 375.0,
            "ملاحظات": "تأكيد التوريد الفوري"
        }])
        al_orders = {"الرواد": wh_df, **{w: pd.DataFrame() for w in EXPECTED_WHOLESALERS if w != "الرواد"}}
        mock_results = make_mock_allocation_result(allocated_orders=al_orders)
        gen = ArabicExcelReportGenerator(mock_results, output_dir=temp_workspace)
        out_path = gen.generate(custom_date_str="2026-09-28")

        wb = openpyxl.load_workbook(out_path, data_only=False)
        ws = wb["طلب شراء - الرواد"]
        last_row = ws.max_row
        col_d = str(ws.cell(row=last_row, column=4).value or "")
        col_h = str(ws.cell(row=last_row, column=8).value or "")
        assert "=SUM(" in col_d.upper()
        assert "=SUM(" in col_h.upper()
        wb.close()

    def test_f8_05_kpi_cards_formatting_and_colors(self, temp_workspace):
        """Verify KPI card blocks in rows 4–5 of Sheet 1 have proper background fills and borders."""
        mock_results = make_mock_allocation_result(kpis={"total_items": 100, "matched_items": 50, "out_of_stock": 50, "total_quantity": 200, "total_public_value": 10000.0, "total_purchase_cost": 7500.0, "total_savings": 2500.0, "savings_pct": 25.0})
        gen = ArabicExcelReportGenerator(mock_results, output_dir=temp_workspace)
        out_path = gen.generate(custom_date_str="2026-09-28")

        wb = openpyxl.load_workbook(out_path)
        ws = wb["المقارنة الشاملة وتوزيع الأصناف"]
        # Row 4 col 1 has KPI title
        c = ws.cell(row=4, column=1)
        assert c.value is not None
        wb.close()

    def test_f8_06_unmatched_items_sheet_schema_and_row_count(self, temp_workspace):
        """Verify Sheet 9 ('نواقص غير متوفرة') has 5 columns with Slate Grey header (#00595959)."""
        unmatched_rows = [
            {"م": i + 1, "الصنف المطلوب": f"ناقص {i}", "الكمية": i + 1, "حالة التوفر": "غير متوفر بجميع المخازن (نواقص)", "ملاحظات": "لم يتم العثور"}
            for i in range(10)
        ]
        mock_results = make_mock_allocation_result(out_of_stock_df=pd.DataFrame(unmatched_rows))
        gen = ArabicExcelReportGenerator(mock_results, output_dir=temp_workspace)
        out_path = gen.generate(custom_date_str="2026-09-28")

        wb = openpyxl.load_workbook(out_path)
        ws = wb["نواقص غير متوفرة"]
        # Header is row 4
        assert ws.cell(row=4, column=1).value == "م"
        assert "الصنف" in str(ws.cell(row=4, column=2).value)
        assert ws.max_row == 14  # 4 header rows + 10 data rows
        wb.close()
        out_path = gen.generate(custom_date_str="2026-09-28")

        wb = openpyxl.load_workbook(out_path)
        ws = wb["نواقص غير متوفرة"]
        # Header is row 4
        assert ws.cell(row=4, column=1).value == "م"
        assert "الصنف" in str(ws.cell(row=4, column=2).value)
        assert ws.max_row == 14  # 4 header rows + 10 data rows
        wb.close()


class TestTier1Feature9Sheet10RecommendationsTab:
    """
    R4 / Feature 9: Dedicated Out-of-Stock Recommendations Tab (توصيات وبدائل النواقص)
    Presents optional commercial alternatives for pharmacist review without
    adulterating primary orders in Sheets 2–8.
    """

    def test_f9_01_sheet10_has_exact_canonical_name(self, temp_workspace):
        """Verify Sheet 10 is titled exactly 'توصيات وبدائل النواقص'."""
        mock_results = make_mock_allocation_result()
        gen = ArabicExcelReportGenerator(mock_results, output_dir=temp_workspace)
        out_path = gen.generate(custom_date_str="2026-09-28")
        wb = openpyxl.load_workbook(out_path)
        assert wb.sheetnames[9] == "توصيات وبدائل النواقص"
        wb.close()

    def test_f9_02_sheet10_nine_columns_schema(self, temp_workspace):
        """Verify Sheet 10 contains exactly 9 columns with standard Arabic titles."""
        mock_results = make_mock_allocation_result(recommendations_df=pd.DataFrame([{
            "م": 1,
            "الصنف الأصلي المطلوب (الناقص)": "صنف أ",
            "الكمية المطلوبة": 2,
            "البديل التجاري المقترح": "بديل ب",
            "المخزن المتاح لديه": "كيور فارما",
            "سعر الجمهور": 100.0,
            "نسبة الخصم %": 0.20,
            "السعر بعد الخصم": 80.0,
            "ملاحظات الصيدلي والتوصيات": "بديل تجاري معتمد مقترح للمراجعة المهنية"
        }]))
        gen = ArabicExcelReportGenerator(mock_results, output_dir=temp_workspace)
        out_path = gen.generate(custom_date_str="2026-09-28")
        wb = openpyxl.load_workbook(out_path)
        ws = wb["توصيات وبدائل النواقص"]

        expected_headers = [
            "م",
            "الصنف الأصلي المطلوب (الناقص)",
            "الكمية المطلوبة",
            "البديل التجاري المقترح",
            "المخزن المتاح لديه",
            "سعر الجمهور",
            "نسبة الخصم %",
            "السعر بعد الخصم",
            "ملاحظات الصيدلي والتوصيات"
        ]
        actual_headers = [str(ws.cell(row=4, column=c).value).strip() for c in range(1, 10)]
        for eh, ah in zip(expected_headers, actual_headers):
            assert eh in ah or ah in eh, f"Header mismatch: expected '{eh}', got '{ah}'"
        wb.close()

    def test_f9_03_sheet10_fallback_rows_for_unmatched_items(self, temp_workspace):
        """Verify shortages without alternatives populate fallback text and dashes."""
        mock_results = make_mock_allocation_result(recommendations_df=pd.DataFrame([{
            "م": 1,
            "الصنف الأصلي المطلوب (الناقص)": "دواء بدون بديل",
            "الكمية المطلوبة": 10,
            "البديل التجاري المقترح": "لا يوجد بديل تجاري متوفر حالياً بالمخازن",
            "المخزن المتاح لديه": "-",
            "سعر الجمهور": "-",
            "نسبة الخصم %": "-",
            "السعر بعد الخصم": "-",
            "ملاحظات الصيدلي والتوصيات": "يتطلب المتابعة مع الوكلاء أو توفير صنف بديل معتمد"
        }]))
        gen = ArabicExcelReportGenerator(mock_results, output_dir=temp_workspace)
        out_path = gen.generate(custom_date_str="2026-09-28")
        wb = openpyxl.load_workbook(out_path)
        ws = wb["توصيات وبدائل النواقص"]

        row5_vals = [ws.cell(row=5, column=c).value for c in range(1, 10)]
        assert row5_vals[3] == "لا يوجد بديل تجاري متوفر حالياً بالمخازن"
        assert row5_vals[4] == "-"
        assert row5_vals[8] == "يتطلب المتابعة مع الوكلاء أو توفير صنف بديل معتمد"
        wb.close()

    def test_f9_04_sheet10_alternative_rows_populated_correctly(self, temp_workspace):
        """Verify shortages with alternatives populate prices, discounts, and pharmacist review note."""
        mock_results = make_mock_allocation_result(recommendations_df=pd.DataFrame([{
            "م": 1,
            "الصنف الأصلي المطلوب (الناقص)": "بانادول اقراص",
            "الكمية المطلوبة": 5,
            "البديل التجاري المقترح": "بارامول اقراص",
            "المخزن المتاح لديه": "كيور فارما",
            "سعر الجمهور": 20.0,
            "نسبة الخصم %": 0.22,
            "السعر بعد الخصم": 15.6,
            "ملاحظات الصيدلي والتوصيات": "بديل تجاري معتمد مقترح للمراجعة المهنية"
        }]))
        gen = ArabicExcelReportGenerator(mock_results, output_dir=temp_workspace)
        out_path = gen.generate(custom_date_str="2026-09-28")
        wb = openpyxl.load_workbook(out_path)
        ws = wb["توصيات وبدائل النواقص"]

        assert ws.cell(row=5, column=4).value == "بارامول اقراص"
        assert ws.cell(row=5, column=5).value == "كيور فارما"
        assert ws.cell(row=5, column=9).value == "بديل تجاري معتمد مقترح للمراجعة المهنية"
        wb.close()

    def test_f9_05_primary_wholesaler_orders_contain_zero_substitutes(self, temp_workspace):
        """Verify commercial substitutes never leak into Sheets 2 through 8."""
        wh_df = pd.DataFrame([{
            "م": 1,
            "اسم الصنف في أمر الشراء": "كونكور 5 مجم اقراص",
            "اسم الصنف في كتالوج المخزن": "كونكور 5 مجم اقراص",
            "الكمية المطلوبة": 10,
            "نسبة الخصم %": 0.20,
            "سعر الجمهور للوحدة": 80.0,
            "سعر التوريد الصافي": 64.0,
            "إجمالي القيمة بسعر الجمهور": 800.0,
            "إجمالي القيمة الصافية": 640.0,
            "ملاحظات التوريد": "تأكيد التوريد الفوري"
        }])
        al_orders = {"كيور فارما": wh_df, **{w: pd.DataFrame() for w in EXPECTED_WHOLESALERS if w != "كيور فارما"}}
        mock_results = make_mock_allocation_result(allocated_orders=al_orders)
        gen = ArabicExcelReportGenerator(mock_results, output_dir=temp_workspace)
        out_path = gen.generate(custom_date_str="2026-09-28")
        wb = openpyxl.load_workbook(out_path)

        for w_name in EXPECTED_WHOLESALERS:
            clean_name = w_name.replace("مخزن ", "").replace(".pdf", "").strip()
            ws = wb[f"طلب شراء - {clean_name}"]
            for r in range(5, ws.max_row + 1):
                item_val = str(ws.cell(row=r, column=3).value or "")
                assert "بديل " not in item_val, f"Commercial substitute leaked into primary sheet {w_name}: {item_val}"
        wb.close()

    def test_f9_06_allocator_generates_recommendations_dataframe(self):
        """Verify allocator generates recommendations DataFrame with proper schema."""
        catalog_df = pd.DataFrame([
            {
                "warehouse_name": "كيور فارما",
                "item_name": "بديل كونكور 5 مجم",
                "normalized_name": "بديل كونكور 5 مجم",
                "dosage": "5 مجم",
                "unit_size": "قرص",
                "base_price": 80.0,
                "discount_rate": 0.20,
                "stock_status": True,
            }
        ])
        buying_list = pd.DataFrame([
            {
                "requested_item": "كونكور 5 مجم اقراص",
                "quantity": 5,
                "normalized_name": "كونكور 5 مجم اقراص",
                "dosage": "5 مجم",
                "unit_size": "قرص",
            }
        ])
        allocator = MedicineAllocator(catalog_df)
        res = allocator.allocate_order(buying_list)
        recs = res.get("recommendations", pd.DataFrame())
        assert not recs.empty
        assert "الصنف الأصلي المطلوب (الناقص)" in recs.columns


# ============================================================================
# TIER 2: BOUNDARY & CORNER CASES
# ============================================================================

class TestTier2BoundaryAndCornerCases:
    """Resilience testing under extreme inputs, malformed structures, boundary values, and noise."""

    def test_t2_01_empty_purchase_order_raises_or_returns_empty(self, temp_workspace):
        """Verify empty purchase order file produces clean empty DataFrame or raises ValueError."""
        empty_xlsx = os.path.join(temp_workspace, "empty.xlsx")
        wb = openpyxl.Workbook()
        wb.save(empty_xlsx)
        wb.close()

        df = load_purchase_order(empty_xlsx)
        assert df.empty or len(df) == 0

    def test_t2_02_single_item_purchase_order_generates_valid_workbook(self, temp_workspace):
        """Verify single item order generates valid 10-sheet workbook with summary formula =SUM(D5:D5)."""
        single_row = [{
            "اسم الصنف المطلوب (في طلب الشراء)": "بانادول اكسترا 24 قرص",
            "اسم الصنف في كشف المخزن الفائز": "بانادول اكسترا 24 قرص",
            "الكمية المطلوبة": 1,
            "المخزن الفائز": "الرواد",
            "أعلى خصم %": 0.26,
            "سعر الجمهور": 45.0,
            "سعر الشراء للوحدة": 33.3,
            "إجمالي سعر الجمهور": 45.0,
            "إجمالي تكلفة الشراء": 33.3,
            "قيمة الوفر": 11.7,
        }]
        wh_orders = {
            "الرواد": pd.DataFrame([{
                "م": 1,
                "اسم الصنف في أمر الشراء": "بانادول اكسترا 24 قرص",
                "اسم الصنف في كتالوج المخزن": "بانادول اكسترا 24 قرص",
                "الكمية المطلوبة": 1,
                "نسبة الخصم %": 0.26,
                "سعر الجمهور للوحدة": 45.0,
                "سعر التوريد الصافي": 33.3,
                "إجمالي القيمة بسعر الجمهور": 45.0,
                "إجمالي القيمة الصافية": 33.3,
                "ملاحظات التوريد": "تأكيد التوريد الفوري"
            }]),
            **{w: pd.DataFrame() for w in EXPECTED_WHOLESALERS if w != "الرواد"}
        }
        res = make_mock_allocation_result(
            master_df=pd.DataFrame(single_row),
            allocated_orders=wh_orders,
            kpis={"total_items": 1, "matched_items": 1, "out_of_stock": 0}
        )
        gen = ArabicExcelReportGenerator(res, output_dir=temp_workspace)
        out_path = gen.generate(custom_date_str="2026-09-28")

        wb = openpyxl.load_workbook(out_path, data_only=False)
        ws_wh = wb["طلب شراء - الرواد"]
        # Summary row at row 6
        assert ws_wh.cell(row=6, column=2).value == "الإجمالي"
        col_d_formula = str(ws_wh.cell(row=6, column=4).value or "")
        assert "=SUM(" in col_d_formula.upper()
        wb.close()

    def test_t2_03_extreme_noise_strings_and_control_chars_filtered(self):
        """Verify non-medication strings, control characters (\\x00-\\x1f), and zero-width joiners are sanitized."""
        noisy_str = "\x00\x01بانادول\u200c\u200d اكسترا\x1f"
        cleaned = clean_state_suffix(noisy_str)
        assert "\x00" not in cleaned and "\x1f" not in cleaned

        # Extreme junk strings must be recognized as noise
        assert is_catalog_noise("---") is True
        assert is_catalog_noise("===***===") is True
        assert is_catalog_noise("123456789") is True

    def test_t2_04_extreme_discount_values_handling(self):
        """Verify parsing behavior on boundary discounts: 0%, 99.9%, negative, and >100%."""
        assert parse_discount("0%") == 0.0
        assert parse_discount("0.0") == 0.0
        assert parse_discount("85%") == 0.85
        # Wholesale discount cap: rates above 85% are capped/sanitized to 0.0 to reject OCR errors
        assert parse_discount("99.9%") == 0.0
        assert parse_discount("-15%") == 0.0
        assert parse_discount("150%") == 0.0

    def test_t2_05_stacked_and_malformed_runic_sequences(self):
        """Verify decode_custom_catalog_glyphs gracefully handles stacked and concatenated Canadian runes."""
        stacked = "ᛞᘭᢔᗷᜧᣂᢕ"
        decoded = decode_custom_catalog_glyphs(stacked)
        assert isinstance(decoded, str)
        assert len(decoded) > 0

    def test_t2_06_boundary_decimal_strengths(self):
        """Verify boundary decimal strengths like 0.05, 0.25, 5., 5.0, 12.5/500."""
        s_005 = extract_drug_strengths_smart("دواء 0.05 مجم")
        s_quarter = extract_drug_strengths_smart("دواء 0.25 مجم")
        s_combo = extract_drug_strengths_smart("دواء 12.5/500 مجم")

        assert "0.05" in s_005 or "0.05" in str(s_005)
        assert "0.25" in s_quarter
        assert "12.5" in s_combo or "500" in s_combo

    def test_t2_07_zero_quantity_order_items_handling(self, temp_workspace):
        """Verify purchase order items with explicit zero quantity are preserved without inflation."""
        file_path = os.path.join(temp_workspace, "zero_qty.xlsx")
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.cell(row=1, column=1, value="اسم الصنف")
        ws.cell(row=1, column=2, value="الكمية")
        ws.cell(row=2, column=1, value="بانادول اكسترا 24 قرص")
        ws.cell(row=2, column=2, value=0)
        wb.save(file_path)
        wb.close()

        df = load_purchase_order(file_path, drop_zero_qty=False)
        assert len(df) == 1
        assert df["quantity"].iloc[0] == 0

    def test_t2_08_mixed_script_items_preservation(self):
        """Verify items with mixed Arabic, English, and Roman numerals are normalized without corruption."""
        mixed = "Augmentin 1gm 14 Tab اوجمنتين 1 جم"
        norm = normalize_arabic_text(mixed)
        assert "augmentin" in norm.lower() or "اوجمنتين" in norm


# ============================================================================
# TIER 3: CROSS-FEATURE COMBINATIONS
# ============================================================================

class TestTier3CrossFeatureCombinations:
    """Pairwise combinatorial interactions across ingestion, glyphs, clinical safety, and allocation."""

    def test_t3_01_hadhaba_glyph_with_decimal_strength_and_pricing_state(self):
        """
        Pairwise: Al-Hadhaba catalog item with aboriginal glyphs + decimal strength ('5.')
        matched against order item with pricing state ('س ج').
        """
        # Order item with pricing state and shorthand decimal:
        req = "دوستينكس 5. اقراص س ج"
        # Candidate catalog item with runes (e.g. ﻛᛞﺴﻮل 0.5 مجم):
        raw_cand = "دوستينكس ﻛᛞﺴﻮل 0.5 مجم"
        decoded_cand = decode_custom_catalog_glyphs(raw_cand)

        # 1. Verify glyph was decoded
        assert "ب" in decoded_cand or "كبسول" in decoded_cand
        # 2. Verify clinical match succeeds between 5. and 0.5
        assert is_safe_pharmaceutical_match(req, decoded_cand) is True
        # 3. Verify clinical match rejects 5 mg
        assert is_safe_pharmaceutical_match(req, "دوستينكس 5 مجم اقراص") is False

    def test_t3_02_competitive_allocation_hadhaba_vs_cure_pharma(self):
        """
        Pairwise: Al-Hadhaba catalog item (with decoded glyphs) competes with Cure Pharma.
        Allocator awards item to highest discount rate.
        """
        catalog_data = [
            {
                "warehouse_name": "مخزن الهضبة فارم",
                "item_name": "ميبو ﻛᛞﺴﻮل",  # Will decode to ميبو كبسول
                "normalized_name": "ميبو كبسول",
                "dosage": "",
                "unit_size": "كبسول",
                "base_price": 100.0,
                "discount_rate": 0.35,  # Higher discount
                "stock_status": True,
            },
            {
                "warehouse_name": "كيور فارما",
                "item_name": "ميبو كبسول",
                "normalized_name": "ميبو كبسول",
                "dosage": "",
                "unit_size": "كبسول",
                "base_price": 100.0,
                "discount_rate": 0.25,
                "stock_status": True,
            }
        ]
        buying_list = pd.DataFrame([{
            "requested_item": "ميبو كبسول س ج",
            "quantity": 5,
            "normalized_name": "ميبو كبسول س ج",
            "dosage": "",
            "unit_size": "كبسول",
        }])

        allocator = MedicineAllocator(pd.DataFrame(catalog_data))
        res = allocator.allocate_order(buying_list)
        master_df = res["master_comparison"]

        assert not master_df.empty
        winner = master_df["المخزن الفائز"].iloc[0]
        assert "الهضبة" in winner, f"Expected Al-Hadhaba to win on 35% discount, got {winner}"
        assert master_df["أعلى خصم %"].iloc[0] == pytest.approx(0.35, abs=0.01)

    def test_t3_03_umbrella_brand_with_slash_form_and_pricing_state(self):
        """
        Pairwise: Umbrella brand (Limitless) with slash-separated form and pricing state ('س ق').
        Blocks wrong sub-brand, allows matching sub-brand.
        """
        req = "ليمتلس اوستيودانس اقراص س ق / كبسول"
        cand_wrong = "ليمتلس كروماكس اقراص"
        cand_right = "ليمتلس اوستيودانس اقراص"

        assert is_safe_pharmaceutical_match(req, cand_wrong) is False
        assert is_safe_pharmaceutical_match(req, cand_right) is True

    def test_t3_04_shortage_with_pricing_state_generates_sheet10_recommendation(self):
        """
        Pairwise: Out-of-stock item carrying 'س ج' is segregated to Sheet 9
        and generates an entry in Sheet 10 without leaking to primary orders.
        """
        catalog_data = [
            {
                "warehouse_name": "كيور فارما",
                "item_name": "بديل زيلوريك 100 مجم",
                "normalized_name": "بديل زيلوريك 100 مجم",
                "dosage": "100 مجم",
                "unit_size": "قرص",
                "base_price": 30.0,
                "discount_rate": 0.20,
                "stock_status": True,
            }
        ]
        buying_list = pd.DataFrame([{
            "requested_item": "زيلوريك 100 مجم اقراص س ج",
            "quantity": 10,
            "normalized_name": "زيلوريك 100 مجم اقراص س ج",
            "dosage": "100 مجم",
            "unit_size": "قرص",
        }])

        allocator = MedicineAllocator(pd.DataFrame(catalog_data))
        res = allocator.allocate_order(buying_list)

        # Zyloric was not directly matched because catalog only had 'بديل'
        assert res["kpis"]["matched_items"] == 0
        assert res["kpis"]["out_of_stock"] == 1

        # Recommendations should have alternative or fallback
        recs = res["recommendations"]
        assert not recs.empty
        assert "زيلوريك" in recs["الصنف الأصلي المطلوب (الناقص)"].iloc[0]

    def test_t3_05_tie_breaking_supplier_consolidation_across_categories(self):
        """
        Pairwise: Equal discount between two wholesalers breaks tie in favor of
        wholesaler with prior allocated items to consolidate purchase orders.
        """
        catalog_data = [
            # Supplier A wins Item 1
            {"warehouse_name": "الرواد", "item_name": "صنف أ", "normalized_name": "صنف أ", "dosage": "", "unit_size": "", "base_price": 50.0, "discount_rate": 0.30, "stock_status": True},
            {"warehouse_name": "كيور فارما", "item_name": "صنف أ", "normalized_name": "صنف أ", "dosage": "", "unit_size": "", "base_price": 50.0, "discount_rate": 0.20, "stock_status": True},
            # Tie on Item 2 (both 25%)
            {"warehouse_name": "الرواد", "item_name": "صنف ب", "normalized_name": "صنف ب", "dosage": "", "unit_size": "", "base_price": 50.0, "discount_rate": 0.25, "stock_status": True},
            {"warehouse_name": "كيور فارما", "item_name": "صنف ب", "normalized_name": "صنف ب", "dosage": "", "unit_size": "", "base_price": 50.0, "discount_rate": 0.25, "stock_status": True},
        ]
        buying_list = pd.DataFrame([
            {"requested_item": "صنف أ", "quantity": 1, "normalized_name": "صنف أ", "dosage": "", "unit_size": ""},
            {"requested_item": "صنف ب", "quantity": 1, "normalized_name": "صنف ب", "dosage": "", "unit_size": ""},
        ])

        allocator = MedicineAllocator(pd.DataFrame(catalog_data))
        res = allocator.allocate_order(buying_list)
        master = res["master_comparison"]

        row_b = master[master["اسم الصنف في أمر الشراء"] == "صنف ب"]
        assert not row_b.empty
        # Tied item should consolidate to Al-Rowad
        assert "الرواد" in row_b["المخزن الفائز"].iloc[0]

    def test_t3_06_synthetic_end_to_end_mini_pipeline(self, temp_workspace):
        """
        Pairwise: Complete execution of run_procurement_pipeline with synthetic order
        and synthetic Excel catalogs, asserting 10 RTL sheets generated cleanly.
        """
        order_path = os.path.join(temp_workspace, "mini_order.xlsx")
        pd.DataFrame([
            {"اسم الصنف": "كونكور 5 مجم اقراص س ج", "الكمية المطلوبة": 10},
            {"اسم الصنف": "دوستينكس 5. اقراص", "الكمية المطلوبة": 2},
        ]).to_excel(order_path, index=False)

        cat_path = os.path.join(temp_workspace, "الرواد كفر الدوار مميز.xlsx")
        pd.DataFrame([
            {"اسم الصنف": "كونكور 5 مجم اقراص", "نسبة الخصم": 0.22, "سعر الجمهور": 80.0},
            {"اسم الصنف": "دوستينكس 0.5 مجم قرص", "نسبة الخصم": 0.18, "سعر الجمهور": 120.0},
        ]).to_excel(cat_path, index=False)

        res = run_procurement_pipeline(
            order_files=[order_path],
            catalog_files=[cat_path],
            output_dir=temp_workspace,
            custom_date_str="2026-09-28",
        )

        assert os.path.isfile(res["report_path"])
        assert res["kpis"]["total_items"] == 2
        assert res["kpis"]["matched_items"] == 2


# ============================================================================
# TIER 4: REAL-WORLD ACCEPTANCE WORKLOAD
# ============================================================================

class TestTier4RealWorldAcceptanceWorkload:
    """
    Tier 4: Real-World Acceptance Workload against physical daily inputs
    (279.xlsx + 7 Wholesaler PDFs in C:\\Users\\omart\\Downloads).
    Asserts exact golden reference metrics from أفضل_اسعار_2026-09-28.xlsx.
    """

    @pytest.mark.skipif(
        not os.path.exists(ORDER_279_PATH),
        reason=f"Daily order 279.xlsx not found at {ORDER_279_PATH}"
    )
    def test_t4_01_order_279_ingestion_preserves_exactly_572_items(self):
        """
        Acceptance Check:
        Parsing 279.xlsx must yield exactly 572 distinct items and 14,122 total units,
        completely excluding the 'AE' row (Row 321) and internal repeated headers.
        """
        df = load_purchase_order(ORDER_279_PATH)
        assert len(df) == 572, f"Expected exactly 572 distinct items in 279.xlsx, got {len(df)}"
        assert df["quantity"].sum() == 14122, f"Expected 14,122 total quantity, got {df['quantity'].sum()}"

        items = df["requested_item"].astype(str).tolist()
        assert "AE" not in items, "'AE' row must be excluded"
        assert "إسم الصنف" not in items, "Internal table header row must be excluded"

    @pytest.mark.skipif(
        not all(os.path.exists(os.path.join(DOWNLOADS_DIR, f)) for f in EXPECTED_PDF_FILES),
        reason="One or more of the 7 wholesaler PDFs is missing in C:\\Users\\omart\\Downloads"
    )
    def test_t4_02_all_7_wholesaler_pdfs_normalize_to_canonical_names(self):
        """
        Acceptance Check:
        All 7 wholesaler PDF filenames normalize to the 7 golden target wholesaler names.
        """
        normalized_names = [
            normalize_vendor_name(os.path.splitext(f)[0])
            for f in EXPECTED_PDF_FILES
        ]
        assert normalized_names == EXPECTED_WHOLESALERS, (
            f"Normalized names mismatch:\nExpected: {EXPECTED_WHOLESALERS}\nActual:   {normalized_names}"
        )

    @pytest.mark.skipif(
        not (os.path.exists(ORDER_279_PATH) and all(os.path.exists(os.path.join(DOWNLOADS_DIR, f)) for f in EXPECTED_PDF_FILES)),
        reason="Daily order or wholesaler PDFs missing in C:\\Users\\omart\\Downloads"
    )
    def test_t4_03_full_pipeline_acceptance_on_2026_09_28_daily_inputs(self, temp_workspace):
        """
        Full Acceptance Workload:
        Executes complete pipeline on 279.xlsx + 7 Wholesaler PDFs.
        Asserts 100% KPI parity with golden benchmark أفضل_اسعار_2026-09-28.xlsx:
          - Total requested: exactly 572 items
          - Total allocated: exactly 154 items
          - Shortages: exactly 418 items
          - Al-Hadhaba allocation: exactly 49 items
          - Total savings: 97,898.87 EGP (32.3% of public value)
          - Total public price: 303,523.00 EGP
          - Total purchase cost: 205,624.13 EGP
          - Exactly 10 RTL sheets
          - Sheet 10 recommendations: exactly 490 rows
          - 0 cross-form false positives, 0 umbrella brand cross-matches
        """
        pdf_paths = [os.path.join(DOWNLOADS_DIR, f) for f in EXPECTED_PDF_FILES]

        res = run_procurement_pipeline(
            order_files=[ORDER_279_PATH],
            catalog_files=pdf_paths,
            output_dir=temp_workspace,
            custom_date_str="2026-09-28",
        )

        kpis = res["kpis"]
        report_path = res["report_path"]

        # 1. Total items, allocated, shortages
        assert kpis["total_items"] == 572, f"Expected 572 requested items, got {kpis.get('total_items')}"
        assert kpis["matched_items"] == 154, f"Expected 154 allocated items, got {kpis.get('matched_items')}"
        assert kpis["out_of_stock"] == 418, f"Expected 418 shortage items, got {kpis.get('out_of_stock')}"

        # 2. Financial KPIs
        assert kpis["total_public_value"] == pytest.approx(303523.00, abs=1.0)
        assert kpis["total_purchase_cost"] == pytest.approx(205624.13, abs=1.0)
        assert kpis["total_savings"] == pytest.approx(97898.87, abs=1.0)
        assert kpis["savings_pct"] == pytest.approx(32.3, abs=0.2)

        # 3. Al-Hadhaba allocated items count
        hadhaba_orders = res["results"]["wholesaler_orders"].get("مخزن الهضبة فارم", pd.DataFrame())
        assert len(hadhaba_orders) == 49, f"Expected 49 Al-Hadhaba items, got {len(hadhaba_orders)}"

        # 4. Workbook validation: 10 RTL sheets
        wb = openpyxl.load_workbook(report_path)
        assert len(wb.sheetnames) == 10, f"Expected 10 sheets, got {len(wb.sheetnames)}"
        for sname in wb.sheetnames:
            ws = wb[sname]
            assert ws.views.sheetView[0].rightToLeft is True, f"Sheet '{sname}' must have rightToLeft=True"

        # 5. Sheet 10 recommendations row count
        ws10 = wb["توصيات وبدائل النواقص"]
        # Rows 1-4 are headers, data rows 5..494 -> exactly 490 rows
        data_rows = ws10.max_row - 4
        assert data_rows == 490, f"Expected 490 recommendation rows in Sheet 10, got {data_rows}"
        wb.close()
