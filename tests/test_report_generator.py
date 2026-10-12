# -*- coding: utf-8 -*-
import os
import tempfile
import pytest
import openpyxl
import pandas as pd
import numpy as np

from src.report_generator import ArabicExcelReportGenerator, clean_excel_str, clean_sheet_title


class TestReportGeneratorSanitization:
    def test_clean_excel_str_removes_control_chars(self):
        user_problem_str = f'1اد 8{chr(6)} 318ا ج'
        cleaned = clean_excel_str(user_problem_str)
        assert chr(6) not in cleaned
        assert cleaned == '1اد 8 318ا ج'

        dirty_str = f'Panadol{chr(0)}{chr(1)}{chr(7)}{chr(8)} Extra{chr(11)}{chr(12)} 500mg{chr(31)}'
        cleaned_dirty = clean_excel_str(dirty_str)
        assert cleaned_dirty == 'Panadol Extra 500mg'

        assert clean_excel_str(None) == ''
        assert clean_excel_str(float('nan')) == ''
        assert clean_excel_str(np.nan) == ''
        assert clean_excel_str('nan') == ''

    def test_clean_sheet_title_removes_forbidden_chars(self):
        forbidden_title = 'مخزن [القاهرة/الجيزة:1*?]'
        cleaned_title = clean_sheet_title(forbidden_title)
        for char in [':', '\\', '/', '?', '*', '[', ']']:
            assert char not in cleaned_title
        assert len(cleaned_title) <= 31

    def test_generate_report_with_illegal_characters_succeeds(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            user_problem_str = f'1اد 8{chr(6)} 318ا ج'
            dirty_cat_item = f'Panadol{chr(1)} Extra {chr(30)} 500mg'
            dirty_wh_name = f'مخزن الإيمان{chr(7)} [فرع:1]'
            dirty_oos_item = f'دواء مفقود{chr(0)} تماماً'
            dirty_note = f'ملاحظة توريد خاصة{chr(31)}'

            master_df = pd.DataFrame([
                {
                    'م': 1,
                    'اسم الصنف في أمر الشراء': user_problem_str,
                    'اسم الصنف في كتالوج المخزن': dirty_cat_item,
                    'الكمية': 5,
                    'المخزن الفائز': dirty_wh_name,
                    'أعلى نسبة خصم': 0.22,
                    'سعر الجمهور': 100.0,
                    'سعر الشراء': 78.0,
                    'إجمالي الجمهور': 500.0,
                    'إجمالي الشراء': 390.0,
                    'الوفر': 110.0,
                    dirty_wh_name: 0.22,
                },
                {
                    'م': 2,
                    'اسم الصنف في أمر الشراء': f'Brufen{chr(11)} 400mg',
                    'اسم الصنف في كتالوج المخزن': f'Brufen{chr(12)} 400mg 30 Tab',
                    'الكمية': 2,
                    'المخزن الفائز': dirty_wh_name,
                    'أعلى نسبة خصم': np.nan,
                    'سعر الجمهور': np.nan,
                    'سعر الشراء': 30.0,
                    'إجمالي الجمهور': 0.0,
                    'إجمالي الشراء': 60.0,
                    'الوفر': 0.0,
                    dirty_wh_name: np.nan,
                }
            ])

            po_df = pd.DataFrame([
                {
                    'م': 1,
                    'اسم الصنف في أمر الشراء': user_problem_str,
                    'اسم الصنف في كتالوج المخزن': dirty_cat_item,
                    'الكمية المطلوبة': 5,
                    'نسبة الخصم %': 0.22,
                    'سعر الجمهور للوحدة': 100.0,
                    'سعر التوريد الصافي': 78.0,
                    'إجمالي القيمة بسعر الجمهور': 500.0,
                    'إجمالي القيمة الصافية': 390.0,
                    'ملاحظات التوريد': dirty_note,
                }
            ])

            oos_df = pd.DataFrame([
                {
                    'م': 1,
                    'اسم الصنف المطلوب': dirty_oos_item,
                    'الكمية المطلوبة': 10,
                    'حالة التوفر': 'غير متوفر',
                    'ملاحظات المتابعة والبحث': f'متابعة{chr(8)} مع الوكيل',
                }
            ])

            allocated_orders = {
                dirty_wh_name: po_df,
                f'{dirty_wh_name} (إضافي)': po_df,
            }

            result_dict = {
                'kpis': {
                    'total_items': 2,
                    'matched_items': 1,
                    'out_of_stock': 1,
                    'total_quantity': 7,
                    'total_public_value': 500.0,
                    'total_purchase_cost': 390.0,
                    'total_savings': 110.0,
                    'savings_pct': 22.0,
                    'num_warehouses': 2,
                },
                'master_comparison': master_df,
                'allocated_orders': allocated_orders,
                'out_of_stock': oos_df,
            }

            generator = ArabicExcelReportGenerator(result_dict, output_dir=tmpdir)
            report_path = generator.generate()

            assert os.path.isfile(report_path)
            assert os.path.getsize(report_path) > 0

            wb = openpyxl.load_workbook(report_path)
            assert 'المقارنة الشاملة وتوزيع الأصناف' in wb.sheetnames
            assert 'نواقص غير متوفرة' in wb.sheetnames
            assert len(wb.sheetnames) >= 4

            ws_master = wb['المقارنة الشاملة وتوزيع الأصناف']
            cell_val = ws_master.cell(row=9, column=2).value
            assert cell_val == '1اد 8 318ا ج'
            assert chr(6) not in cell_val
