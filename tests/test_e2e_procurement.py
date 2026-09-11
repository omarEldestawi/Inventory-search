# -*- coding: utf-8 -*-
"""
E2E Test Suite: Egyptian Pharmaceutical Procurement Engine (Tiers 3 & 4)
Opaque-box verification against domain rules, clinical safety gates,
and multi-sheet Right-to-Left Arabic Excel report specifications.

Coverage:
  - Tier 3: Pairwise Combinations (PDF+XLSX, multi-PDF+CSV, Excel catalogs, mixed batches, competitive allocation)
  - Tier 4: Real-world Workloads, Arabic RTL verification, Sheet Naming, Formula Validation, Clinical Safety Gates, High-Volume Scaling
"""
import os
import sys
import glob
import tempfile
import pytest
import pandas as pd
import openpyxl

from src.pipeline import run_procurement_pipeline
from src.allocator import MedicineAllocator


# ============================================================================
# FIXTURES
# ============================================================================

@pytest.fixture
def sample_catalog_pdfs():
    """Returns list of the three benchmark sample catalog PDFs."""
    base_dir = os.path.join(os.path.dirname(__file__), "..", "input", "catalogs")
    files = sorted([
        os.path.abspath(os.path.join(base_dir, f))
        for f in os.listdir(base_dir)
        if f.endswith(".pdf")
    ])
    assert len(files) >= 3, f"Expected 3 catalog PDFs in {base_dir}, found {len(files)}"
    return files


@pytest.fixture
def primary_order_xlsx():
    """Returns absolute path to primary sample order (طلبية_اليوم.xlsx)."""
    p = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "input", "orders", "طلبية_اليوم.xlsx"))
    assert os.path.exists(p), f"Missing sample order: {p}"
    return p


@pytest.fixture
def arabic_order_xlsx():
    """Returns absolute path to Arabic test order (طلبية_عربي_اختبار.xlsx)."""
    p = os.path.abspath(os.path.join(os.path.dirname(__file__), "..", "input", "orders", "طلبية_عربي_اختبار.xlsx"))
    assert os.path.exists(p), f"Missing Arabic sample order: {p}"
    return p


@pytest.fixture
def temp_output_dir():
    """Provides isolated temporary directory for test report outputs."""
    with tempfile.TemporaryDirectory() as tmpdir:
        yield tmpdir


# ============================================================================
# TIER 3: PAIRWISE COMBINATORIAL TESTS
# ============================================================================

class TestTier3PairwiseCombinations:
    """Pairwise combinations of file formats, catalog sources, and allocation rules."""

    def test_pairwise_pdf_catalog_and_xlsx_order(self, sample_catalog_pdfs, primary_order_xlsx, temp_output_dir):
        """Tier 3: PDF catalog + XLSX order pairwise combination."""
        res = run_procurement_pipeline(
            order_files=[primary_order_xlsx],
            catalog_files=[sample_catalog_pdfs[0]],
            output_dir=temp_output_dir,
        )
        assert os.path.isfile(res["report_path"])
        assert res["kpis"]["total_items"] == 8

    def test_pairwise_multi_pdf_and_csv_order(self, sample_catalog_pdfs, temp_output_dir):
        """Tier 3: Multi-PDF catalogs + CSV order pairwise combination."""
        csv_order_path = os.path.join(temp_output_dir, "order.csv")
        csv_data = [
            {"اسم الصنف": "Panadol Extra 24 Tab", "الكمية المطلوبة": 10},
            {"اسم الصنف": "Augmentin 1gm 14 Tab", "الكمية المطلوبة": 5},
            {"اسم الصنف": "Brufen 400mg 30 Tab", "الكمية المطلوبة": 20},
        ]
        pd.DataFrame(csv_data).to_csv(csv_order_path, index=False, encoding="utf-8-sig")

        res = run_procurement_pipeline(
            order_files=[csv_order_path],
            catalog_files=sample_catalog_pdfs,
            output_dir=temp_output_dir,
        )

        kpis = res["kpis"]
        assert kpis["total_items"] == 3
        assert kpis["matched_items"] == 3
        assert kpis["out_of_stock"] == 0

    def test_pairwise_excel_catalog_and_xlsx_order(self, primary_order_xlsx, temp_output_dir):
        """Tier 3: Excel catalog (.xlsx) + XLSX order pairwise combination."""
        excel_catalog_path = os.path.join(temp_output_dir, "wholesaler_al_nour.xlsx")
        cat_items = [
            {"الصنف": "Panadol Extra 24 Tab", "الخصم": "26.0%", "السعر": 45.0},
            {"الصنف": "Augmentin 1gm 14 Tab", "الخصم": "27.5%", "السعر": 120.0},
            {"الصنف": "Cataflam 50mg 20 Tab", "الخصم": "22.0%", "السعر": 55.0},
            {"الصنف": "Concor 5mg 30 Tab", "الخصم": "19.0%", "السعر": 80.0},
            {"الصنف": "Brufen 400mg 30 Tab", "الخصم": "17.0%", "السعر": 60.0},
            {"الصنف": "Nexium 40mg 28 Tab", "الخصم": "29.0%", "السعر": 280.0},
            {"الصنف": "Flagyl 500mg 20 Tab", "الخصم": "24.0%", "السعر": 30.0},
        ]
        pd.DataFrame(cat_items).to_excel(excel_catalog_path, index=False)

        res = run_procurement_pipeline(
            order_files=[primary_order_xlsx],
            catalog_files=[excel_catalog_path],
            output_dir=temp_output_dir,
        )

        kpis = res["kpis"]
        assert kpis["total_items"] == 8
        assert kpis["matched_items"] == 7
        assert kpis["out_of_stock"] == 1  # Amaryl not in catalog
        assert kpis["avg_discount"] > 20.0

    def test_pairwise_mixed_catalogs_pdf_excel_csv(self, sample_catalog_pdfs, primary_order_xlsx, temp_output_dir):
        """Tier 3: Heterogeneous catalogs (PDF + XLSX + CSV) combined simultaneously."""
        xlsx_cat = os.path.join(temp_output_dir, "extra_vendor.xlsx")
        csv_cat = os.path.join(temp_output_dir, "discount_vendor.csv")

        pd.DataFrame([
            {"اسم الصنف": "Amaryl 2mg 30 Tab", "نسبة الخصم": 0.22, "سعر الجمهور": 50.0}
        ]).to_excel(xlsx_cat, index=False)

        pd.DataFrame([
            {"اسم الصنف": "Panadol Extra 24 Tab", "نسبة الخصم": "30%", "سعر الجمهور": 45.0}
        ]).to_csv(csv_cat, index=False, encoding="utf-8-sig")

        mixed_catalogs = sample_catalog_pdfs + [xlsx_cat, csv_cat]

        res = run_procurement_pipeline(
            order_files=[primary_order_xlsx],
            catalog_files=mixed_catalogs,
            output_dir=temp_output_dir,
        )

        kpis = res["kpis"]
        # Amaryl is now supplied by extra_vendor, so all 8 should match!
        assert kpis["total_items"] == 8
        assert kpis["matched_items"] == 8
        assert kpis["out_of_stock"] == 0

    def test_pairwise_competitive_allocation_max_discount(self, temp_output_dir):
        """Tier 3: Verifies that allocator deterministically awards item to wholesaler with highest discount."""
        order_path = os.path.join(temp_output_dir, "single_order.xlsx")
        pd.DataFrame([{"اسم الصنف": "Brufen 400mg 30 Tab", "الكمية المطلوبة": 10}]).to_excel(order_path, index=False)

        cat_a = os.path.join(temp_output_dir, "مخزن_الوفاء.xlsx")
        cat_b = os.path.join(temp_output_dir, "مخزن_البركة.xlsx")

        pd.DataFrame([{"اسم الصنف": "Brufen 400mg 30 Tab", "نسبة الخصم": 0.15, "سعر الجمهور": 60.0}]).to_excel(cat_a, index=False)
        pd.DataFrame([{"اسم الصنف": "Brufen 400mg 30 Tab", "نسبة الخصم": 0.23, "سعر الجمهور": 60.0}]).to_excel(cat_b, index=False)

        res = run_procurement_pipeline(
            order_files=[order_path],
            catalog_files=[cat_a, cat_b],
            output_dir=temp_output_dir,
        )

        master_df = res["results"]["master_comparison"]
        winner = master_df["المخزن الفائز"].iloc[0]
        winning_disc = master_df["أعلى خصم %"].iloc[0]

        assert "البركة" in winner
        assert winning_disc == pytest.approx(0.23, abs=0.01)

    def test_pairwise_tie_breaking_supplier_consolidation(self, temp_output_dir):
        """Tier 3: Tie-breaking favors supplier with higher current allocation count to minimize order split."""
        order_path = os.path.join(temp_output_dir, "tie_order.xlsx")
        pd.DataFrame([
            {"اسم الصنف": "Panadol Extra 24 Tab", "الكمية المطلوبة": 5},
            {"اسم الصنف": "Cataflam 50mg 20 Tab", "الكمية المطلوبة": 5},
        ]).to_excel(order_path, index=False)

        # Supplier 1 gives 25% on Panadol and 20% on Cataflam
        # Supplier 2 gives 20% on Panadol and 20% on Cataflam (tie on Cataflam)
        cat1 = os.path.join(temp_output_dir, "مخزن_أ.xlsx")
        cat2 = os.path.join(temp_output_dir, "مخزن_ب.xlsx")

        pd.DataFrame([
            {"اسم الصنف": "Panadol Extra 24 Tab", "نسبة الخصم": 0.25, "سعر الجمهور": 45.0},
            {"اسم الصنف": "Cataflam 50mg 20 Tab", "نسبة الخصم": 0.20, "سعر الجمهور": 55.0},
        ]).to_excel(cat1, index=False)

        pd.DataFrame([
            {"اسم الصنف": "Panadol Extra 24 Tab", "نسبة الخصم": 0.20, "سعر الجمهور": 45.0},
            {"اسم الصنف": "Cataflam 50mg 20 Tab", "نسبة الخصم": 0.20, "سعر الجمهور": 55.0},
        ]).to_excel(cat2, index=False)

        res = run_procurement_pipeline(
            order_files=[order_path],
            catalog_files=[cat1, cat2],
            output_dir=temp_output_dir,
        )

        master_df = res["results"]["master_comparison"]
        # Both items should consolidate to مخزن_أ because it already won Panadol
        for _, row in master_df.iterrows():
            assert "مخزن أ" in row["المخزن الفائز"] or "مخزن_أ" in row["المخزن الفائز"]


# ============================================================================
# TIER 4: REAL-WORLD WORKLOADS & ARABIC MULTI-SHEET EXCEL AUDIT
# ============================================================================

class TestTier4RealWorldProcurementWorkloads:
    """Real-world scenarios with benchmark catalogs, Arabic RTL layout, formulas, and clinical safety."""

    def test_real_world_standard_order_execution_and_excel_audit(
        self, sample_catalog_pdfs, primary_order_xlsx, temp_output_dir
    ):
        """
        Tier 4: Complete real-world procurement optimization run using standard benchmark catalogs
        and daily order, with full audit of RTL presentation, sheet names, formulas, and KPIs.
        """
        res = run_procurement_pipeline(
            order_files=[primary_order_xlsx],
            catalog_files=sample_catalog_pdfs,
            output_dir=temp_output_dir,
            custom_date_str="2026-09-11",
        )

        # 1. Output File Verification
        report_path = res["report_path"]
        arabic_path = res["arabic_report_path"]
        assert os.path.isfile(report_path), f"Report not found: {report_path}"
        assert os.path.isfile(arabic_path), f"Arabic alias not found: {arabic_path}"
        assert os.path.getsize(report_path) > 5000, "Workbook file size suspiciously small"

        # 2. KPI Verification
        kpis = res["kpis"]
        assert kpis["total_items"] == 8
        assert kpis["matched_items"] == 7
        assert kpis["out_of_stock"] == 1
        assert kpis["avg_discount"] > 20.0
        assert kpis["total_savings"] == pytest.approx(
            kpis["total_public_value"] - kpis["total_purchase_cost"], abs=0.5
        )
        assert kpis["savings_pct"] > 0
        assert kpis["num_warehouses"] >= 2

        # 3. Load OpenPyXL Workbook for Structural & Layout Audit
        wb = openpyxl.load_workbook(report_path)
        sheet_names = wb.sheetnames

        # Verify Mandatory Sheet Names
        assert "المقارنة الشاملة وتوزيع الأصناف" in sheet_names, "Missing master comparison sheet"
        assert "نواقص غير متوفرة" in sheet_names, "Missing out of stock sheet"
        po_sheets = [s for s in sheet_names if s.startswith("طلب شراء - ")]
        assert len(po_sheets) >= 2, f"Expected at least 2 supplier PO sheets, found {len(po_sheets)}"

        # 4. Right-to-Left (RTL) View Audit for EVERY sheet
        for sheet_name in sheet_names:
            ws = wb[sheet_name]
            # Check sheetView[0].rightToLeft
            is_rtl = False
            try:
                is_rtl = ws.views.sheetView[0].rightToLeft
            except Exception:
                try:
                    is_rtl = ws.sheet_view.rightToLeft
                except Exception:
                    pass
            assert is_rtl is True, f"Sheet '{sheet_name}' is not configured as Right-to-Left (RTL)"

        # 5. Formula Audit in Master Comparison Sheet
        ws_master = wb["المقارنة الشاملة وتوزيع الأصناف"]
        # Search for =SUM formulas in summary row
        found_sum_formulas = []
        for row in ws_master.iter_rows(values_only=False):
            for cell in row:
                if cell.value and isinstance(cell.value, str) and cell.value.startswith("=SUM("):
                    found_sum_formulas.append((cell.coordinate, cell.value))

        assert len(found_sum_formulas) >= 3, (
            f"Expected at least 3 '=SUM()' accounting formulas in master sheet, found: {found_sum_formulas}"
        )

        # 6. Formula Audit in Supplier Purchase Order Sheets
        for po_sheet_name in po_sheets:
            ws_po = wb[po_sheet_name]
            po_sums = []
            for row in ws_po.iter_rows(values_only=False):
                for cell in row:
                    if cell.value and isinstance(cell.value, str) and cell.value.startswith("=SUM("):
                        po_sums.append((cell.coordinate, cell.value))
            assert len(po_sums) >= 2, (
                f"Expected '=SUM()' formulas in PO sheet '{po_sheet_name}', found: {po_sums}"
            )

        # 7. Out of Stock Sheet Audit
        ws_oos = wb["نواقص غير متوفرة"]
        oos_items = []
        for r in range(5, ws_oos.max_row + 1):
            val = ws_oos.cell(row=r, column=2).value
            if val:
                oos_items.append(str(val).strip())
        assert any("Amaryl" in item for item in oos_items), (
            f"Expected 'Amaryl 2mg 30 Tab' in Out of Stock sheet, found: {oos_items}"
        )

    def test_real_world_arabic_order_workload(self, sample_catalog_pdfs, arabic_order_xlsx, temp_output_dir):
        """Tier 4: Ingestion and optimization of pure Arabic daily order with Arabic trade names."""
        res = run_procurement_pipeline(
            order_files=[arabic_order_xlsx],
            catalog_files=sample_catalog_pdfs,
            output_dir=temp_output_dir,
        )

        kpis = res["kpis"]
        assert kpis["total_items"] > 0
        assert kpis["matched_items"] > 0

        # Verify workbook generated
        wb = openpyxl.load_workbook(res["report_path"])
        ws = wb["المقارنة الشاملة وتوزيع الأصناف"]
        assert ws.views.sheetView[0].rightToLeft is True

    def test_real_world_clinical_safety_gate_enforcement(self, temp_output_dir):
        """
        Tier 4: Clinical Safety Verification.
        Enforces that hard safety gates (Form, Strength, Modifier, Infant Milk Stage) strictly
        prevent dangerous clinical mismatches and route them to Out of Stock.
        """
        order_path = os.path.join(temp_output_dir, "dangerous_order.xlsx")
        order_items = [
            {"اسم الصنف": "Augmentin 1gm 14 Tab", "الكمية المطلوبة": 10},  # Strength mismatch vs 625mg
            {"اسم الصنف": "Antodine 20mg Tab", "الكمية المطلوبة": 5},      # Form mismatch vs Amp
            {"اسم الصنف": "Panadol Extra", "الكمية المطلوبة": 8},          # Modifier mismatch vs Cold & Flu
            {"اسم الصنف": "Hero Baby 1 Milk", "الكمية المطلوبة": 4},       # Stage mismatch vs Stage 2
        ]
        pd.DataFrame(order_items).to_excel(order_path, index=False)

        cat_path = os.path.join(temp_output_dir, "trap_catalog.xlsx")
        cat_items = [
            {"اسم الصنف": "Augmentin 625mg 14 Tab", "نسبة الخصم": "35%", "سعر الجمهور": 90.0},
            {"اسم الصنف": "Antodine 20mg 3 Amp", "نسبة الخصم": "30%", "سعر الجمهور": 40.0},
            {"اسم الصنف": "Panadol Cold & Flu 24 Tab", "نسبة الخصم": "32%", "سعر الجمهور": 45.0},
            {"اسم الصنف": "Hero Baby 2 Milk 400g", "نسبة الخصم": "25%", "سعر الجمهور": 150.0},
        ]
        pd.DataFrame(cat_items).to_excel(cat_path, index=False)

        res = run_procurement_pipeline(
            order_files=[order_path],
            catalog_files=[cat_path],
            output_dir=temp_output_dir,
        )

        kpis = res["kpis"]
        # ALL 4 dangerous items must be REJECTED by clinical safety gates
        assert kpis["total_items"] == 4
        assert kpis["matched_items"] == 0, f"Dangerous matches made: {res['results']['master_comparison']}"
        assert kpis["out_of_stock"] == 4

        # Verify they are all listed in Out of Stock
        oos_df = res["results"]["out_of_stock"]
        assert len(oos_df) == 4

    def test_real_world_large_order_scalability(self, sample_catalog_pdfs, temp_output_dir):
        """Tier 4: Scalability test with 50+ item order exercising performance, memory, and formatting."""
        large_order_path = os.path.join(temp_output_dir, "large_order.xlsx")
        
        # Build 50 diverse items combining matches and unmatches
        base_items = [
            "Panadol Extra 24 Tab",
            "Augmentin 1gm 14 Tab",
            "Cataflam 50mg 20 Tab",
            "Concor 5mg 30 Tab",
            "Brufen 400mg 30 Tab",
            "Nexium 40mg 28 Tab",
            "Flagyl 500mg 20 Tab",
            "Amaryl 2mg 30 Tab",
            "Voltaren 100mg 10 Supp",
            "Congestal 20 Tab",
        ]
        order_rows = []
        for i in range(50):
            item_name = base_items[i % len(base_items)]
            if i >= 40:
                item_name = f"RareMedication_{i} 50mg Tab"
            order_rows.append({"اسم الصنف": item_name, "الكمية المطلوبة": (i % 10) + 1})

        pd.DataFrame(order_rows).to_excel(large_order_path, index=False)

        res = run_procurement_pipeline(
            order_files=[large_order_path],
            catalog_files=sample_catalog_pdfs,
            output_dir=temp_output_dir,
        )

        kpis = res["kpis"]
        assert kpis["total_items"] == len(pd.DataFrame(order_rows)["اسم الصنف"].unique())
        assert kpis["matched_items"] > 0
        assert os.path.isfile(res["report_path"])

    def test_real_world_excel_visual_styling_parity(self, sample_catalog_pdfs, primary_order_xlsx, temp_output_dir):
        """Tier 4: Visual and typographical styling audit: Navy headers, white bold font, zebra fills."""
        res = run_procurement_pipeline(
            order_files=[primary_order_xlsx],
            catalog_files=sample_catalog_pdfs,
            output_dir=temp_output_dir,
        )

        wb = openpyxl.load_workbook(res["report_path"])
        ws = wb["المقارنة الشاملة وتوزيع الأصناف"]

        # Row 8 is the table header row
        header_cell = ws.cell(row=8, column=2)
        assert header_cell.font.bold is True
        # Check navy blue fill (1F4E79)
        assert header_cell.fill.start_color.rgb in ["001F4E79", "1F4E79"]
        assert header_cell.font.color.rgb in ["00FFFFFF", "FFFFFF"]
