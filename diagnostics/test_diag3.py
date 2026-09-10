# -*- coding: utf-8 -*-
import sys
import openpyxl
import unicodedata

sys.stdout.reconfigure(encoding='utf-8')

bench_path = r'C:\Users\omart\Downloads\Procurement_Order_Optimized_2026-09-10.xlsx'
our_path = r'G:\My Drive\Pharama Daily orders\Procurement_Order_Optimized_2026-09-10.xlsx'

wb_bench = openpyxl.load_workbook(bench_path)
ws_bench = wb_bench['المقارنة الشاملة وتوزيع الأصناف']
bench_rows = list(ws_bench.iter_rows(values_only=True))

bench_items = []
for r in bench_rows[6:]:
    if r[0] is not None and isinstance(r[0], int):
        bench_items.append({
            'num': r[0],
            'req_name': str(r[1]).strip() if r[1] else '',
            'cat_name': str(r[2]).strip() if r[2] else '',
            'qty': r[3],
            'winner': str(r[4]).strip() if r[4] else '',
            'discount': str(r[5]).strip() if r[5] else '',
            'public_price': r[6],
            'buy_price': r[7]
        })

wb_our = openpyxl.load_workbook(our_path)
ws_our = wb_our['المقارنة الشاملة وتوزيع الأصناف']
our_rows = list(ws_our.iter_rows(values_only=True))

our_items = []
for r in our_rows[8:]:
    if r[0] is not None and isinstance(r[0], int):
        our_items.append({
            'num': r[0],
            'req_name': str(r[1]).strip() if r[1] else '',
            'cat_name': str(r[2]).strip() if r[2] else '',
            'qty': r[3],
            'winner': str(r[4]).strip() if r[4] else '',
            'discount': r[5],
            'public_price': r[6],
            'buy_price': r[7]
        })

def norm(s):
    if not s: return ''
    s = unicodedata.normalize('NFKD', s)
    s = ''.join(c for c in s if not unicodedata.combining(c))
    return ' '.join(s.replace('  ', ' ').strip().lower().split())

our_by_req = {norm(x['req_name']): x for x in our_items}
bench_by_req = {norm(x['req_name']): x for x in bench_items}

print("=== 10 ITEMS IN OUR OUTPUT BUT NOT IN BENCHMARK ===")
count = 0
for o in our_items:
    if norm(o['req_name']) not in bench_by_req:
        count += 1
        print(f"{count}. Req: '{o['req_name']}'")
        print(f"   Our Matched Cat: '{o['cat_name']}' | Winner: '{o['winner']}' | Disc: {o['discount']} | PubPrice: {o['public_price']}")

print("\n=== 31 ITEMS IN BENCHMARK BUT NOT IN OUR OUTPUT ===")
count = 0
for b in bench_items:
    if norm(b['req_name']) not in our_by_req:
        count += 1
        print(f"{count}. [Bench #{b['num']}] Req: '{b['req_name']}'")
        print(f"   Bench Matched Cat: '{b['cat_name']}' | Winner: '{b['winner']}' | Disc: {b['discount']} | PubPrice: {b['public_price']}")

print("\n=== 23 ITEMS IN BOTH ===")
count = 0
for b in bench_items:
    if norm(b['req_name']) in our_by_req:
        count += 1
        o = our_by_req[norm(b['req_name'])]
        same_winner = (b['winner'] in o['winner'] or o['winner'] in b['winner'])
        print(f"{count}. Req: '{b['req_name']}'")
        print(f"   Bench: Cat='{b['cat_name']}', Winner='{b['winner']}', Disc={b['discount']}")
        print(f"   Ours : Cat='{o['cat_name']}', Winner='{o['winner']}', Disc={o['discount']}")
        print(f"   Winner Agreement: {same_winner}")

