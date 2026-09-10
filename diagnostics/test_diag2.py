# -*- coding: utf-8 -*-
import sys
import openpyxl

sys.stdout.reconfigure(encoding='utf-8')

bench_path = r'C:\Users\omart\Downloads\Procurement_Order_Optimized_2026-09-10.xlsx'
our_path = r'G:\My Drive\Pharama Daily orders\Procurement_Order_Optimized_2026-09-10.xlsx'

wb_bench = openpyxl.load_workbook(bench_path)
ws_bench = wb_bench['المقارنة الشاملة وتوزيع الأصناف']
bench_rows = list(ws_bench.iter_rows(values_only=True))

bench_items = []
for r in bench_rows[6:]:
    if r[0] is not None and isinstance(r[0], int):
        bench_items.append({
            'num': r[0],
            'req_name': str(r[1]).strip() if r[1] else '',
            'cat_name': str(r[2]).strip() if r[2] else '',
            'qty': r[3],
            'winner': str(r[4]).strip() if r[4] else '',
            'discount': str(r[5]).strip() if r[5] else '',
            'public_price': r[6],
            'buy_price': r[7]
        })

wb_our = openpyxl.load_workbook(our_path)
ws_our = wb_our['المقارنة الشاملة وتوزيع الأصناف']
our_rows = list(ws_our.iter_rows(values_only=True))

our_items = []
for r in our_rows[8:]:
    if r[0] is not None and isinstance(r[0], int):
        our_items.append({
            'num': r[0],
            'req_name': str(r[1]).strip() if r[1] else '',
            'cat_name': str(r[2]).strip() if r[2] else '',
            'qty': r[3],
            'winner': str(r[4]).strip() if r[4] else '',
            'discount': r[5],
            'public_price': r[6],
            'buy_price': r[7]
        })

print(f"Total Benchmark items: {len(bench_items)}")
print(f"Total Our items: {len(our_items)}")

# Let's inspect our items
print("\n--- OUR 33 ITEMS ---")
for o in our_items:
    print(f"#{o['num']} Req: '{o['req_name']}' | Cat: '{o['cat_name']}' | Winner: '{o['winner']}' | Disc: {o['discount']}")

# Let's check fuzzy or normalized match between bench req_name and our req_name
import unicodedata
def norm(s):
    if not s: return ''
    s = unicodedata.normalize('NFKD', s)
    s = ''.join(c for c in s if not unicodedata.combining(c))
    return ' '.join(s.replace('  ', ' ').strip().lower().split())

our_norm_map = {norm(o['req_name']): o for o in our_items}
bench_norm_map = {norm(b['req_name']): b for b in bench_items}

matched = []
missing_from_ours = []
for b in bench_items:
    nb = norm(b['req_name'])
    if nb in our_norm_map:
        matched.append((b, our_norm_map[nb]))
    else:
        # Check partial
        found = None
        for ko, vo in our_norm_map.items():
            if ko == nb or ko in nb or nb in ko:
                found = vo
                break
        if found:
            matched.append((b, found))
        else:
            missing_from_ours.append(b)

print(f"\nMatched count: {len(matched)}")
print(f"Missing count: {len(missing_from_ours)}")

