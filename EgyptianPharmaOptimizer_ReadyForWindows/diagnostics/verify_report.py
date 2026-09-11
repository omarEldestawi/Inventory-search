# -*- coding: utf-8 -*-
import sys
import openpyxl

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook('output/أفضل_اسعار_2026-09-01.xlsx')
ws = wb['Master Comparison']
rows = [r for r in ws.iter_rows(values_only=True) if r and r[0] is not None and isinstance(r[0], int)]
print(f"Total matched items in Master Comparison: {len(rows)}")

req_items = [r[1] for r in rows]
dup_reqs = [it for it in set(req_items) if req_items.count(it) > 1]
print(f"Duplicate rows in Master Comparison: {len(dup_reqs)}")

ws_oos = wb['Out of Stock - النواقص']
oos_rows = [r for r in ws_oos.iter_rows(values_only=True) if r and r[0] is not None and isinstance(r[0], int)]
print(f"Total items in Out of Stock: {len(oos_rows)}")
print(f"Total unique items overall: {len(rows)} + {len(oos_rows)} = {len(rows) + len(oos_rows)}")

print("\nSample 25 verified allocations:")
for r in rows[:25]:
    print(f"#{r[0]:>2}: Req='{r[1]}' (Qty={r[2]}) -> Winner='{r[4]}' (Disc={r[5]}) | InWH='{r[6]}'")
