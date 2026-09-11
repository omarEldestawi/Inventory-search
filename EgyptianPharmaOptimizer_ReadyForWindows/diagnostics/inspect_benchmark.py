# -*- coding: utf-8 -*-
import sys
import os
import openpyxl
import pandas as pd

if hasattr(sys.stdout, 'reconfigure'):
    sys.stdout.reconfigure(encoding='utf-8')

wb = openpyxl.load_workbook('input/Procurement_Order_Optimized.xlsx')
print("Sheets in Procurement_Order_Optimized.xlsx:", wb.sheetnames)

ws = wb['Master Comparison']
rows = list(ws.iter_rows(values_only=True))
print("Header Row 7:", rows[6])
print("\nFirst 10 matched items in Master Comparison:")
for r in rows[7:17]:
    if any(r):
        print(f"Item #{r[0]}: Req='{r[1]}' | Form='{r[3]}' | Winner='{r[4]}' | Disc='{r[5]}' | InWH='{r[6]}'")

print(f"\nTotal matched items in template: {len([r for r in rows[7:] if any(r) and r[0] is not None])}")
